from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from reconstruct import (
    inspect_source,
    build_packet as _build_packet,
    verify_packet as _verify_packet,
    mapping_sha256,
)


def build_packet(source: Path, mapping: dict, destination: Path) -> dict:
    approved_hash = mapping_sha256(mapping)
    report = _build_packet(
        source, mapping, destination,
        approved_mapping_sha256=approved_hash,
        approved_by="Synthetic test reviewer",
    )
    receipt = destination.parent / f"{destination.name}-approval.json"
    receipt.write_text(json.dumps({
        "approved_by": "Synthetic test reviewer",
        "mapping_sha256": approved_hash,
    }, indent=2) + "\n", encoding="utf-8")
    return report


def verify_packet(source: Path, destination: Path) -> dict:
    receipt = destination.parent / f"{destination.name}-approval.json"
    approval = json.loads(receipt.read_text(encoding="utf-8"))
    return _verify_packet(source, destination, approved_mapping_sha256=approval["mapping_sha256"])


def create_fixture(root: Path) -> tuple[Path, dict]:
    source = root / "synthetic-screen.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Layout"
    for column in range(1, 13):
        sheet.column_dimensions[chr(64 + column)].width = 3
    sheet.merge_cells("B2:L2")
    sheet["B2"] = "Synthetic application"
    sheet.merge_cells("B4:D4")
    sheet["B4"] = "Applicant"
    sheet["F4"] = "FLD-001"
    sheet["H4"] = "Text"
    sheet["J4"] = "0010"
    sheet["B7"] = "Amount"
    sheet["F7"] = "FLD-002"
    sheet["H7"] = "Decimal"
    sheet["B10"] = "Only when amount exceeds 100, confirmation is required."
    sheet["B11"] = "Do not register when confirmation is declined."
    sheet["B14"] = "Unmapped annotation"
    sheet["B16"] = "=SUM(1,2)"
    sheet["L12"] = "Hidden column annotation"
    sheet["N20"] = "Outside print area annotation"
    sheet.row_dimensions[14].hidden = True
    sheet.column_dimensions["L"].hidden = True
    sheet.print_area = "A1:L18"
    book.save(source)
    book.close()
    inventory = inspect_source(source, ["Layout"])
    mapping = {
        "schema_version": 1,
        "source_sha256": inventory["source_sha256"],
        "sheets": ["Layout"],
        "fields": [
            {"field_id": ["Layout!F4"], "name": ["Layout!B4"],
             "data_type": ["Layout!H4"], "length": ["Layout!J4"]},
            {"field_id": ["Layout!F7"], "name": ["Layout!B7"],
             "data_type": ["Layout!H7"]},
        ],
        "rules": [{"condition": ["Layout!B10"], "action": ["Layout!B11"]}],
        "excluded": {"Layout!B2": "Document title; retained in source inventory."},
    }
    return source, mapping


def create_generic_fixture(root: Path) -> tuple[Path, dict]:
    source = root / "synthetic-report.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Report"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Quarterly summary"
    sheet["A3"] = "Revenue"
    sheet["B3"] = "1200"
    sheet["A4"] = "Returns"
    sheet["B4"] = "30"
    sheet["A7"] = "Product code"
    sheet["B7"] = "Quantity"
    sheet["A8"] = "P-001"
    sheet["B8"] = "4"
    sheet.print_area = "A1:B8"
    book.save(source)
    book.close()
    inventory = inspect_source(source, ["Report"])
    mapping = {
        "schema_version": 2,
        "source_sha256": inventory["source_sha256"],
        "sheets": ["Report"],
        "tables": [
            {
                "name": "Metrics",
                "columns": ["metric name", "amount"],
                "rows": [
                    {"metric name": ["Report!A3"], "amount": ["Report!B3"]},
                    {"metric name": ["Report!A4"], "amount": ["Report!B4"]},
                ],
            },
            {
                "name": "Line items",
                "columns": ["product_code", "quantity"],
                "rows": [{"product_code": ["Report!A8"], "quantity": ["Report!B8"]}],
            },
        ],
        "excluded": {"Report!A1": "Synthetic report title; retained in source inventory."},
    }
    return source, mapping


class ReconstructionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source, self.mapping = create_fixture(self.root)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_rejects_output_package_tampering_with_refreshed_manifest(self) -> None:
        for alteration in ("hyperlink", "extra_part", "table_style"):
            with self.subTest(alteration=alteration):
                destination = self.root / alteration
                build_packet(self.source, self.mapping, destination)
                output = destination / "reconstructed.xlsx"
                if alteration == "hyperlink":
                    book = load_workbook(output)
                    book["Fields"]["B2"].hyperlink = "https://example.invalid/"
                    book.save(output)
                    book.close()
                elif alteration == "table_style":
                    book = load_workbook(output)
                    next(iter(book["Fields"].tables.values())).tableStyleInfo.name = "TableStyleMedium3"
                    book.save(output)
                    book.close()
                else:
                    with ZipFile(output, "a") as archive:
                        archive.writestr("xl/extra.xml", "<extra/>")
                manifest_path = destination / "manifest.json"
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                manifest["artifact_hashes"]["reconstructed.xlsx"] = hashlib.sha256(output.read_bytes()).hexdigest()
                manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "output package"):
                    verify_packet(self.source, destination)

    def test_preserves_raw_date_and_numeric_storage(self) -> None:
        book = load_workbook(self.source)
        book.iso_dates = True
        sheet = book["Layout"]
        sheet["A21"] = 45292.25
        sheet["A21"].number_format = "yyyy-mm-dd hh:mm"
        sheet["A22"] = 0.5
        sheet["A22"].number_format = "hh:mm"
        sheet["A23"] = 1.25
        sheet["A23"].number_format = "[h]:mm"
        sheet["A24"] = datetime(2024, 1, 1, 6, 0)
        book.save(self.source)
        book.close()
        inventory = inspect_source(self.source, ["Layout"])
        records = {record["source_ref"]: record for record in inventory["records"]}
        for reference, expected in (("Layout!A21", "45292.25"), ("Layout!A22", "0.5"), ("Layout!A23", "1.25")):
            self.assertEqual(records[reference]["value"], expected)
            self.assertEqual(records[reference]["source_type"], "n")
        self.assertEqual(records["Layout!A24"]["value"], "2024-01-01T06:00:00")
        self.assertEqual(records["Layout!A24"]["source_type"], "d")
        self.mapping["source_sha256"] = inventory["source_sha256"]
        self.mapping["fields"][0]["default"] = ["Layout!A21"]
        destination = self.root / "date-values"
        build_packet(self.source, self.mapping, destination)
        output = load_workbook(destination / "reconstructed.xlsx")
        try:
            self.assertEqual(output["Fields"]["G2"].value, "45292.25")
        finally:
            output.close()

    def test_rejects_columns_beyond_excel_limit_before_generation(self) -> None:
        source, mapping = create_generic_fixture(self.root)
        mapping["tables"][0]["columns"] = [f"column-{index}" for index in range(16384)]
        mapping["tables"][0]["rows"] = []
        destination = self.root / "too-wide"
        with self.assertRaisesRegex(ValueError, "column limit"):
            build_packet(source, mapping, destination)
        self.assertFalse(destination.exists())

    def test_preserves_values_traceability_and_source(self) -> None:
        original = self.source.read_bytes()
        destination = self.root / "packet"
        build_packet(self.source, self.mapping, destination)
        report = verify_packet(self.source, destination)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(report["status"], "draft_with_unresolved")
        self.assertEqual(report["font_name"], "Noto Sans JP")
        self.assertEqual(report["approval"]["approved_by"], "Synthetic test reviewer")
        self.assertEqual(report["approval"]["mapping_sha256"], mapping_sha256(self.mapping))
        receipt = destination.parent / "packet-approval.json"
        self.assertTrue(receipt.is_file())
        self.assertFalse(receipt.is_relative_to(destination))
        book = load_workbook(destination / "reconstructed.xlsx")
        try:
            self.assertEqual(book.sheetnames, ["Fields", "Rules", "SourceMap", "SourceInventory", "Unresolved"])
            fields = list(book["Fields"].values)
            self.assertEqual(fields[1][1:5], ("FLD-001", "Applicant", "Text", "0010"))
            self.assertIsNone(fields[1][5])
            self.assertIn("Do not register when confirmation is declined.", list(book["Rules"].values)[1])
            references = list(book["SourceMap"].values)
            self.assertTrue(any("Layout!B4" in row for row in references))
            source_inventory = list(book["SourceInventory"].values)
            inventory_columns = source_inventory[0]
            inventory_rows = {
                row[inventory_columns.index("source_ref")]: dict(zip(inventory_columns, row))
                for row in source_inventory[1:]
            }
            self.assertTrue(inventory_rows["Layout!B14"]["hidden_row"])
            self.assertTrue(inventory_rows["Layout!L12"]["hidden_column"])
            self.assertEqual(inventory_rows["Layout!N20"]["inside_print_area"], "False")
            unresolved = list(book["Unresolved"].values)
            self.assertTrue(any("Layout!B14" in row for row in unresolved))
            self.assertTrue(any("Layout!N20" in row for row in unresolved))
            self.assertTrue(any("Layout!B16" in row for row in unresolved))
            for sheet in book:
                self.assertFalse(sheet.merged_cells.ranges)
                self.assertEqual(len(sheet.tables), 1)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        finally:
            book.close()

    def test_preserves_english_source_values_and_approved_headers(self) -> None:
        source, mapping = create_generic_fixture(self.root)
        mapping["font_name"] = "Aptos"
        destination = self.root / "english-packet"
        report = build_packet(source, mapping, destination)
        verify_packet(source, destination)
        self.assertEqual(report["font_name"], "Aptos")

        book = load_workbook(destination / "reconstructed.xlsx")
        try:
            metrics = list(book["Metrics"].values)
            self.assertEqual(metrics[0], ("row_key", "metric name", "amount"))
            self.assertEqual(metrics[1][1:], ("Revenue", "1200"))
            self.assertEqual(metrics[2][1:], ("Returns", "30"))
            items = list(book["Line items"].values)
            self.assertEqual(items[0], ("row_key", "product_code", "quantity"))
            self.assertEqual(items[1][1:], ("P-001", "4"))
            self.assertEqual(book["Metrics"]["A1"].font.name, "Aptos")
            inventory_values = {
                str(value) for row in book["SourceInventory"].values for value in row
            }
            self.assertIn("Quarterly summary", inventory_values)
        finally:
            book.close()

    def test_rejects_changed_source_missing_anchor_and_unknown_columns(self) -> None:
        for update, message in [
            ({"source_sha256": "0" * 64}, "hash"),
            ({"fields": [{"name": ["Layout!C4"]}]}, "reference"),
            ({"fields": [{"invented": ["Layout!B4"]}]}, "column"),
        ]:
            mapping = {**self.mapping, **update}
            with self.subTest(message=message):
                with self.assertRaisesRegex(ValueError, message):
                    build_packet(self.source, mapping, self.root / message)
                self.assertFalse((self.root / message).exists())

    def test_rejects_formulas_as_semantic_values_and_unreasoned_exclusion(self) -> None:
        mapping = copy.deepcopy(self.mapping)
        mapping["fields"][0]["default"] = ["Layout!B16"]
        with self.assertRaisesRegex(ValueError, "formula"):
            build_packet(self.source, mapping, self.root / "formula")
        mapping = copy.deepcopy(self.mapping)
        mapping["excluded"]["Layout!B14"] = ""
        with self.assertRaisesRegex(ValueError, "reason"):
            build_packet(self.source, mapping, self.root / "exclusion")

    def test_retains_drawing_package_as_unresolved(self) -> None:
        book = load_workbook(self.source)
        sheet = book.active
        sheet["K17"] = 1
        sheet["K18"] = 2
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=11, min_row=17, max_row=18))
        sheet.add_chart(chart, "L1")
        book.save(self.source)
        book.close()
        self.mapping["source_sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        build_packet(self.source, self.mapping, self.root / "drawings")
        report = verify_packet(self.source, self.root / "drawings")
        self.assertGreater(report["unresolved_count"], 0)
        book = load_workbook(self.root / "drawings/reconstructed.xlsx")
        try:
            self.assertTrue(any("drawing" in str(row).lower() for row in book["Unresolved"].values))
        finally:
            book.close()

    def test_retains_unrecognized_package_parts_as_unresolved(self) -> None:
        with ZipFile(self.source, "a", compression=ZIP_DEFLATED) as archive:
            archive.writestr(
                "xl/externalLinks/externalLink1.xml",
                "<externalLink><externalBook r:id='rId1'/></externalLink>",
            )
        inventory = inspect_source(self.source, ["Layout"])
        references = {record["source_ref"] for record in inventory["records"]}
        self.assertNotIn("PART:_rels/.rels", references)
        part = next((record for record in inventory["records"]
                     if record["source_ref"] == "PART:xl/externalLinks/externalLink1.xml"), None)
        self.assertIsNotNone(part)
        self.mapping["source_sha256"] = inventory["source_sha256"]
        build_packet(self.source, self.mapping, self.root / "external-link")
        packet = load_workbook(self.root / "external-link/reconstructed.xlsx")
        try:
            unresolved = list(packet["Unresolved"].values)
            self.assertTrue(any("externalLink1.xml" in str(row) for row in unresolved))
        finally:
            packet.close()

    def test_refuses_overwrite_and_detects_artifact_changes(self) -> None:
        destination = self.root / "packet"
        build_packet(self.source, self.mapping, destination)
        with self.assertRaises(FileExistsError):
            build_packet(self.source, self.mapping, destination)
        output = destination / "reconstructed.xlsx"
        book = load_workbook(output)
        book["Fields"]["C2"] = "Changed"
        book.save(output)
        book.close()
        with self.assertRaisesRegex(ValueError, "hash"):
            verify_packet(self.source, destination)

    def test_detects_value_changes_even_when_output_hash_is_updated(self) -> None:
        destination = self.root / "packet"
        build_packet(self.source, self.mapping, destination)
        output = destination / "reconstructed.xlsx"
        book = load_workbook(output)
        book["Fields"]["C2"] = "Wrong field"
        book.save(output)
        book.close()
        manifest_path = destination / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["artifact_hashes"]["reconstructed.xlsx"] = hashlib.sha256(output.read_bytes()).hexdigest()
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "values differ"):
            verify_packet(self.source, destination)

    def test_rejects_coordinated_mapping_and_workbook_replacement(self) -> None:
        approved_mapping = copy.deepcopy(self.mapping)
        approved_hash = mapping_sha256(approved_mapping)
        destination = self.root / "approved-packet"
        build_packet(self.source, approved_mapping, destination)

        changed_mapping = copy.deepcopy(approved_mapping)
        changed_mapping["fields"][0]["name"] = ["Layout!B7"]
        replacement = self.root / "replacement-packet"
        build_packet(self.source, changed_mapping, replacement)
        for filename in ["mapping.json", "reconstructed.xlsx", "manifest.json"]:
            shutil.copyfile(replacement / filename, destination / filename)

        with self.assertRaisesRegex(ValueError, "approved mapping hash mismatch"):
            verify_packet(self.source, destination)

    def test_rejects_build_without_explicit_matching_approval(self) -> None:
        with self.assertRaisesRegex(ValueError, "approved mapping hash mismatch"):
            _build_packet(
                self.source, self.mapping, self.root / "unapproved",
                approved_mapping_sha256="0" * 64,
                approved_by="Synthetic test reviewer",
            )
        with self.assertRaisesRegex(ValueError, "approval record name"):
            _build_packet(
                self.source, self.mapping, self.root / "unnamed-approval",
                approved_mapping_sha256=mapping_sha256(self.mapping),
                approved_by=" ",
            )

    def test_cli_fingerprint_build_and_verify_share_approved_hash(self) -> None:
        self.mapping["font_name"] = "Aptos"
        mapping_path = self.root / "mapping.json"
        mapping_path.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding="utf-8")
        script = Path(__file__).resolve().parents[1] / "scripts" / "reconstruct.py"
        common = ["--workbook", str(self.source), "--mapping", str(mapping_path)]
        fingerprint = subprocess.run(
            [sys.executable, str(script), "fingerprint", *common],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(fingerprint.returncode, 0, fingerprint.stderr)
        self.assertIn(f"MAPPING_REVIEW_READY {mapping_sha256(self.mapping)}", fingerprint.stdout)
        self.assertIn('"font_name": "Aptos"', fingerprint.stdout)

        packet = self.root / "cli-packet"
        approval = ["--approved-mapping-sha256", mapping_sha256(self.mapping),
                    "--approved-by", "Synthetic test reviewer"]
        build = subprocess.run(
            [sys.executable, str(script), "build", *common, *approval,
             "--font-name", "Aptos", "--out-dir", str(packet)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(build.returncode, 0, build.stderr + build.stdout)
        self.assertIn('"font_name": "Aptos"', build.stdout)
        verify = subprocess.run(
            [sys.executable, str(script), "verify", "--workbook", str(self.source),
             "--packet", str(packet), *approval[:2]],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(verify.returncode, 0, verify.stderr + verify.stdout)
        workbook = load_workbook(packet / "reconstructed.xlsx")
        try:
            self.assertTrue(all(cell.font.name == "Aptos" for sheet in workbook for row in sheet for cell in row))
        finally:
            workbook.close()
        tampered = subprocess.run(
            [sys.executable, str(script), "verify", "--workbook", str(self.source),
             "--packet", str(packet), "--approved-mapping-sha256", "0" * 64],
            capture_output=True, text=True, check=False,
        )
        self.assertNotEqual(tampered.returncode, 0)
        self.assertIn("approved mapping hash mismatch", tampered.stdout)

    def test_font_name_is_recorded_and_verified(self) -> None:
        mapping = copy.deepcopy(self.mapping)
        mapping["font_name"] = "Aptos"
        destination = self.root / "font-packet"
        approved_hash = mapping_sha256(mapping)
        report = _build_packet(
            self.source, mapping, destination,
            approved_mapping_sha256=approved_hash,
            approved_by="Synthetic test reviewer",
        )
        self.assertEqual(report["font_name"], "Aptos")
        output = destination / "reconstructed.xlsx"
        book = load_workbook(output)
        try:
            book["Fields"]["A1"].font = Font(name="Calibri")
            book.save(output)
        finally:
            book.close()
        manifest_path = destination / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["artifact_hashes"]["reconstructed.xlsx"] = hashlib.sha256(output.read_bytes()).hexdigest()
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "font mismatch"):
            _verify_packet(self.source, destination, approved_mapping_sha256=approved_hash)

    def test_rejects_manifest_and_workbook_font_change_after_approval(self) -> None:
        mapping = copy.deepcopy(self.mapping)
        mapping["font_name"] = "Aptos"
        approved_hash = mapping_sha256(mapping)
        destination = self.root / "tampered-font-packet"
        _build_packet(
            self.source, mapping, destination,
            approved_mapping_sha256=approved_hash,
            approved_by="Synthetic test reviewer",
        )
        output = destination / "reconstructed.xlsx"
        book = load_workbook(output)
        try:
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        cell.font = Font(name="Calibri", size=11)
            book.save(output)
        finally:
            book.close()
        manifest_path = destination / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["font_name"] = "Calibri"
        manifest["artifact_hashes"]["reconstructed.xlsx"] = hashlib.sha256(output.read_bytes()).hexdigest()
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "manifest font name mismatch"):
            _verify_packet(self.source, destination, approved_mapping_sha256=approved_hash)

    def test_rejects_font_override_not_in_approved_mapping(self) -> None:
        approved_hash = mapping_sha256(self.mapping)
        destination = self.root / "unapproved-font-packet"
        with self.assertRaisesRegex(ValueError, "font name must match approved mapping"):
            _build_packet(
                self.source, self.mapping, destination,
                approved_mapping_sha256=approved_hash,
                approved_by="Synthetic test reviewer",
                font_name="Aptos",
            )
        self.assertFalse(destination.exists())

    def test_legacy_manifest_defaults_to_noto_sans_jp(self) -> None:
        destination = self.root / "legacy-font-packet"
        approved_hash = mapping_sha256(self.mapping)
        _build_packet(
            self.source, self.mapping, destination,
            approved_mapping_sha256=approved_hash,
            approved_by="Synthetic test reviewer",
        )
        manifest_path = destination / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        del manifest["font_name"]
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        report = _verify_packet(self.source, destination, approved_mapping_sha256=approved_hash)
        self.assertEqual(report["font_name"], "Noto Sans JP")

    def test_reconstructs_generic_tables_with_cell_traceability(self) -> None:
        source, mapping = create_generic_fixture(self.root)
        mapping_path = self.root / "generic-mapping.json"
        mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        script = Path(__file__).resolve().parents[1] / "scripts" / "reconstruct.py"
        fingerprint = subprocess.run(
            [sys.executable, str(script), "fingerprint", "--workbook", str(source), "--mapping", str(mapping_path)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(fingerprint.returncode, 0, fingerprint.stderr)
        self.assertIn('"Metrics": 2', fingerprint.stdout)
        self.assertIn('"Line items": 1', fingerprint.stdout)
        destination = self.root / "generic-packet"
        approved_hash = mapping_sha256(mapping)
        build = subprocess.run(
            [sys.executable, str(script), "build", "--workbook", str(source), "--mapping", str(mapping_path),
             "--approved-mapping-sha256", approved_hash, "--approved-by", "Synthetic test reviewer",
             "--out-dir", str(destination)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(build.returncode, 0, build.stderr + build.stdout)
        report = _verify_packet(source, destination, approved_mapping_sha256=approved_hash)
        self.assertEqual(report["status"], "draft_with_unresolved")
        self.assertEqual(report["row_counts"]["Metrics"], 2)
        self.assertEqual(report["row_counts"]["Line items"], 1)

        book = load_workbook(destination / "reconstructed.xlsx")
        try:
            self.assertEqual(book.sheetnames, ["Metrics", "Line items", "SourceMap", "SourceInventory", "Unresolved"])
            self.assertEqual(list(book["Metrics"].values)[1], ("Metrics-0001", "Revenue", "1200"))
            self.assertEqual(list(book["Line items"].values)[1], ("Line items-0001", "P-001", "4"))
            source_map = list(book["SourceMap"].values)
            self.assertTrue(any(row[0] == "Metrics" and row[2] == "metric name" and row[3] == "Report!A3"
                                for row in source_map[1:]))
            self.assertTrue(any(row[0] == "Line items" and row[3] == "Report!B8" for row in source_map[1:]))
            table_names = [table.displayName for sheet in book for table in sheet.tables.values()]
            self.assertEqual(len(table_names), len(set(table_names)))
            self.assertTrue(all(" " not in name for name in table_names))
            self.assertTrue(any("LAYOUT:Report" in row for row in book["Unresolved"].values))
        finally:
            book.close()

    def test_rejects_invalid_generic_table_definitions(self) -> None:
        source, original = create_generic_fixture(self.root)
        invalid_mappings = [
            ({"tables": [original["tables"][0], {**original["tables"][1], "name": "Metrics"}]}, "duplicate table name"),
            ({"tables": [{**original["tables"][0], "name": "SourceMap"}, original["tables"][1]]}, "reserved table name"),
            ({"tables": [{**original["tables"][0], "columns": ["metric name", "Metric name"]}, original["tables"][1]]}, "unique"),
        ]
        for update, message in invalid_mappings:
            mapping = {**original, **update}
            destination = self.root / message.replace(" ", "-")
            with self.subTest(message=message):
                with self.assertRaisesRegex(ValueError, message):
                    build_packet(source, mapping, destination)
                self.assertFalse(destination.exists())

    def test_rejects_overlong_or_forbidden_generic_table_names_and_row_key(self) -> None:
        source, original = create_generic_fixture(self.root)
        invalid_tables = [
            ({**original["tables"][0], "name": "T" * 32}, "invalid or reserved table name"),
            ({**original["tables"][0], "name": "Bad/Name"}, "invalid or reserved table name"),
            ({**original["tables"][0], "name": "History"}, "invalid or reserved table name"),
            ({**original["tables"][0], "name": "'Metrics"}, "invalid or reserved table name"),
            ({**original["tables"][0], "columns": ["row_key"]}, "unique and cannot use row_key"),
            ({**original["tables"][0], "columns": ["C" * 32768], "rows": []}, "losslessly"),
            ({**original["tables"][0], "columns": ["bad\x00column"], "rows": []}, "losslessly"),
        ]
        for index, (invalid_table, message) in enumerate(invalid_tables):
            mapping = copy.deepcopy(original)
            mapping["tables"][0] = invalid_table
            destination = self.root / f"invalid-generic-definition-{index}"
            with self.subTest(message=message):
                with self.assertRaisesRegex(ValueError, message):
                    build_packet(source, mapping, destination)
                self.assertFalse(destination.exists())

    def test_records_and_verifies_openpyxl_version_in_manifest(self) -> None:
        from openpyxl import __version__ as openpyxl_version
        destination = self.root / "version-packet"
        approved_hash = mapping_sha256(self.mapping)
        _build_packet(
            self.source, self.mapping, destination,
            approved_mapping_sha256=approved_hash,
            approved_by="Synthetic test reviewer",
        )
        manifest_path = destination / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["openpyxl_version"], openpyxl_version)
        manifest["openpyxl_version"] = "0.0.0"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "openpyxl version mismatch"):
            _verify_packet(self.source, destination, approved_mapping_sha256=approved_hash)

    def test_generic_cli_build_and_verify_accepts_empty_table(self) -> None:
        source, mapping = create_generic_fixture(self.root)
        mapping["tables"].append({"name": "Empty view", "columns": ["label"], "rows": []})
        mapping_path = self.root / "empty-table-mapping.json"
        mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        script = Path(__file__).resolve().parents[1] / "scripts" / "reconstruct.py"
        approved_hash = mapping_sha256(mapping)
        packet = self.root / "empty-table-packet"
        build = subprocess.run(
            [sys.executable, str(script), "build", "--workbook", str(source), "--mapping", str(mapping_path),
             "--approved-mapping-sha256", approved_hash, "--approved-by", "Synthetic test reviewer",
             "--out-dir", str(packet)],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(build.returncode, 0, build.stderr + build.stdout)
        verify = subprocess.run(
            [sys.executable, str(script), "verify", "--workbook", str(source), "--packet", str(packet),
             "--approved-mapping-sha256", approved_hash],
            capture_output=True, text=True, check=False,
        )
        self.assertEqual(verify.returncode, 0, verify.stderr + verify.stdout)
        report = json.loads(verify.stdout.split(" ", 1)[1])
        self.assertEqual(report["row_counts"]["Empty view"], 0)

    def test_literal_formula_text_does_not_execute_and_long_text_is_rejected(self) -> None:
        book = load_workbook(self.source)
        sheet = book.active
        sheet["B4"] = "=1+1"
        sheet["B4"].data_type = "s"
        book.save(self.source)
        book.close()
        self.mapping["source_sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        build_packet(self.source, self.mapping, self.root / "literal")
        book = load_workbook(self.root / "literal/reconstructed.xlsx")
        try:
            self.assertEqual(book["Fields"]["C2"].value, "=1+1")
            self.assertEqual(book["Fields"]["C2"].data_type, "s")
        finally:
            book.close()
        from reconstruct import validate_text
        with self.assertRaisesRegex(ValueError, "losslessly"):
            validate_text("A" * 32768)

    def test_completed_mapping_still_requires_semantic_review(self) -> None:
        inventory = inspect_source(self.source, ["Layout"])
        for record in inventory["records"]:
            if (record["source_ref"] in {"Layout!B14", "Layout!B16", "Layout!L12",
                                          "Layout!N20", "LAYOUT:Layout"}
                    or record["kind"] == "package_object"):
                self.mapping["excluded"][record["source_ref"]] = "Synthetic test control: reviewed and out of scope."
        report = build_packet(self.source, self.mapping, self.root / "mapped")
        self.assertEqual(report["unresolved_count"], 0)
        self.assertEqual(report["status"], "mapped_pending_semantic_review")

    def test_unknown_sheets_and_legacy_formats_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "unknown sheet"):
            inspect_source(self.source, ["Absent"])
        with self.assertRaisesRegex(ValueError, "xlsx"):
            inspect_source(self.source.with_suffix(".xls"), ["Layout"])

    def test_stored_cell_limit_is_checked_before_openpyxl_load(self) -> None:
        with patch("reconstruct.MAX_STORED_CELLS_PER_SHEET", 1), patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(ValueError, "one million stored cells"):
                inspect_source(self.source, ["Layout"])
        load.assert_not_called()

    def test_unselected_sheet_cell_limit_is_checked_before_openpyxl_load(self) -> None:
        source = self.root / "unselected-large-sheet.xlsx"
        book = Workbook()
        book.active.title = "Selected"
        book.active["A1"] = "Selected value"
        unselected = book.create_sheet("Unselected")
        unselected["A1"] = "First value"
        unselected["B1"] = "Second value"
        book.save(source)
        book.close()
        with patch("reconstruct.MAX_STORED_CELLS_PER_SHEET", 1), patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(ValueError, "one million stored cells: Unselected"):
                inspect_source(source, ["Selected"])
        load.assert_not_called()

    def test_doctype_is_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "doctype-original.xlsx"
        source = self.root / "doctype-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    marker = b"<worksheet"
                    position = content.find(marker)
                    self.assertGreaterEqual(position, 0)
                    doctype = b'<!DOCTYPE worksheet [<!ENTITY expansion "blocked">]>'
                    content = content[:position] + doctype + content[position:]
                target_archive.writestr(member, content)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_metadata_doctypes_are_rejected_before_openpyxl_load(self) -> None:
        for member_name in [
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/styles.xml",
        ]:
            with self.subTest(member=member_name):
                original = self.root / "metadata-doctype-original.xlsx"
                source = self.root / "metadata-doctype-input.xlsx"
                book = Workbook()
                book.active.title = "Layout"
                book.active["A1"] = "value"
                book.save(original)
                book.close()
                with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
                    for member in source_archive.infolist():
                        content = source_archive.read(member.filename)
                        if member.filename == member_name:
                            marker = {
                                "xl/workbook.xml": b"<workbook",
                                "xl/_rels/workbook.xml.rels": b"<Relationships",
                                "xl/styles.xml": b"<styleSheet",
                            }[member_name]
                            position = content.find(marker)
                            self.assertGreaterEqual(position, 0)
                            doctype = b'<!DOCTYPE workbook [<!ENTITY expansion "blocked">]>'
                            content = content[:position] + doctype + content[position:]
                        target_archive.writestr(member, content)
                with patch("reconstruct.load_workbook") as load:
                    with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                        inspect_source(source, ["Layout"])
                load.assert_not_called()

    def test_chart_and_table_doctypes_are_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "chart-table-doctype-original.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Layout"
        sheet["A1"] = "Value"
        sheet["A2"] = 1
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2), titles_from_data=True)
        sheet.add_chart(chart, "C1")
        table = Table(displayName="FixtureTable", ref="A1:A2")
        sheet.add_table(table)
        book.save(original)
        book.close()
        with ZipFile(original) as archive:
            guarded_parts = {"xl/charts/chart1.xml", "xl/tables/table1.xml"}
            self.assertTrue(guarded_parts <= set(archive.namelist()))
            for member_name in guarded_parts:
                with self.subTest(member=member_name):
                    source = self.root / f"{Path(member_name).stem}-doctype.xlsx"
                    with ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
                        for member in archive.infolist():
                            content = archive.read(member.filename)
                            if member.filename == member_name:
                                root_tag = b"<chartSpace" if "charts" in member_name else b"<table"
                                position = content.find(root_tag)
                                self.assertGreaterEqual(position, 0)
                                content = content[:position] + b'<!DOCTYPE part [<!ENTITY expansion "blocked">]>' + content[position:]
                            target_archive.writestr(member.filename, content, compress_type=ZIP_DEFLATED)
                    with patch("reconstruct.load_workbook") as load:
                        with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                            inspect_source(source, ["Layout"])
                    load.assert_not_called()

    def test_nonstandard_worksheet_relationship_is_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "nonstandard-relationship-original.xlsx"
        source = self.root / "nonstandard-relationship-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                if member.filename == "xl/_rels/workbook.xml.rels":
                    relationship = next(
                        part for part in content.split(b"<Relationship") if b"/worksheet" in part
                    )
                    target = re.search(rb'Target="([^"]+)"', relationship)
                    self.assertIsNotNone(target)
                    content = content.replace(target.group(0), b'Target="../custom.xml"', 1)
                target_archive.writestr(member, content)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(ValueError, "unsupported worksheet package path"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_duplicate_relationship_id_is_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "duplicate-relationship-original.xlsx"
        source = self.root / "duplicate-relationship-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                if member.filename == "xl/_rels/workbook.xml.rels":
                    worksheet_relationship = next(
                        b"<Relationship" + part.split(b"/>", 1)[0] + b"/>"
                        for part in content.split(b"<Relationship")[1:]
                        if b"/worksheet" in part
                    )
                    content = content.replace(
                        b"</Relationships>", worksheet_relationship + b"</Relationships>",
                    )
                target_archive.writestr(member, content)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(ValueError, "duplicate workbook relationship id"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_nested_worksheet_doctype_is_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "nested-worksheet-original.xlsx"
        source = self.root / "nested-worksheet-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                name = member.filename
                if name == "xl/worksheets/sheet1.xml":
                    marker = b"<worksheet"
                    position = content.find(marker)
                    self.assertGreaterEqual(position, 0)
                    doctype = b'<!DOCTYPE worksheet [<!ENTITY expansion "blocked">]>'
                    content = content[:position] + doctype + content[position:]
                    name = "xl/worksheets/nested/sheet1.xml"
                elif name == "xl/_rels/workbook.xml.rels":
                    content, replacements = re.subn(
                        rb'Target="/?xl/worksheets/sheet1\.xml"',
                        b'Target="/xl/worksheets/nested/sheet1.xml"',
                        content,
                    )
                    self.assertEqual(replacements, 1)
                target_archive.writestr(name, content)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_nested_worksheet_is_not_duplicated_as_package_part(self) -> None:
        original = self.root / "nested-worksheet-inventory-original.xlsx"
        source = self.root / "nested-worksheet-inventory.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                name = member.filename
                if name == "xl/worksheets/sheet1.xml":
                    name = "xl/worksheets/nested/sheet1.xml"
                elif name == "xl/_rels/workbook.xml.rels":
                    content, replacements = re.subn(
                        rb'Target="/?xl/worksheets/sheet1\.xml"',
                        b'Target="/xl/worksheets/nested/sheet1.xml"',
                        content,
                    )
                    self.assertEqual(replacements, 1)
                target_archive.writestr(name, content)
        inventory = inspect_source(source, ["Layout"])
        references = {record["source_ref"] for record in inventory["records"]}
        self.assertIn("Layout!A1", references)
        self.assertNotIn("PART:xl/worksheets/nested/sheet1.xml", references)

    def test_orphan_standard_worksheet_part_is_retained_in_inventory(self) -> None:
        original = self.root / "orphan-worksheet-original.xlsx"
        source = self.root / "orphan-worksheet-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                target_archive.writestr(member.filename, source_archive.read(member.filename), compress_type=ZIP_DEFLATED)
            target_archive.writestr(
                "xl/worksheets/sheet99.xml",
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
                compress_type=ZIP_DEFLATED,
            )
        inventory = inspect_source(source, ["Layout"])
        references = {record["source_ref"] for record in inventory["records"]}
        self.assertIn("PART:xl/worksheets/sheet99.xml", references)

    def test_nested_chart_relationship_target_is_preflighted(self) -> None:
        original = self.root / "nested-chart-original.xlsx"
        source = self.root / "nested-chart-input.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Layout"
        sheet["A1"] = "Value"
        sheet["A2"] = 1
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2), titles_from_data=True)
        sheet.add_chart(chart, "C1")
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                name = member.filename
                if name == "xl/charts/chart1.xml":
                    position = content.find(b"<chartSpace")
                    self.assertGreaterEqual(position, 0)
                    content = content[:position] + b'<!DOCTYPE chart [<!ENTITY expansion "blocked">]>' + content[position:]
                    name = "xl/vendor/chart1.xml"
                elif name == "xl/drawings/_rels/drawing1.xml.rels":
                    content, replacements = re.subn(
                        rb'Target="([^\"]*chart1\.xml)"',
                        b'Target="../vendor/chart1.xml"',
                        content,
                    )
                    self.assertEqual(replacements, 1)
                target_archive.writestr(name, content, compress_type=ZIP_DEFLATED)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_custom_app_properties_relationship_target_is_preflighted(self) -> None:
        original = self.root / "custom-app-properties-original.xlsx"
        source = self.root / "custom-app-properties-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                name = member.filename
                if name == "docProps/app.xml":
                    position = content.find(b"<Properties")
                    self.assertGreaterEqual(position, 0)
                    content = content[:position] + b'<!DOCTYPE Properties [<!ENTITY expansion "blocked">]>' + content[position:]
                    name = "custom/app.xml"
                elif name == "_rels/.rels":
                    content, replacements = re.subn(
                        rb'Target="docProps/app\.xml"',
                        b'Target="custom/app.xml"',
                        content,
                    )
                    self.assertEqual(replacements, 1)
                target_archive.writestr(name, content, compress_type=ZIP_DEFLATED)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(Exception, "DOCTYPE declarations are not supported"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_nonstandard_worksheet_relationship_type_is_rejected_before_openpyxl_load(self) -> None:
        original = self.root / "relationship-type-original.xlsx"
        source = self.root / "relationship-type-input.xlsx"
        book = Workbook()
        book.active.title = "Layout"
        book.active["A1"] = "value"
        book.save(original)
        book.close()
        with ZipFile(original) as source_archive, ZipFile(source, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)
                if member.filename == "xl/_rels/workbook.xml.rels":
                    content = content.replace(
                        b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
                        b"https://example.invalid/officeDocument/2006/relationships/worksheet",
                    )
                target_archive.writestr(member, content)
        with patch("reconstruct.load_workbook") as load:
            with self.assertRaisesRegex(ValueError, "unsupported sheet relationship"):
                inspect_source(source, ["Layout"])
        load.assert_not_called()

    def test_preserves_array_and_data_table_formulas_and_spaced_disjoint_print_areas(self) -> None:
        from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
        source = self.root / "formulas-and-print-area.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Screen Layout"
        sheet["A1"] = "Field"
        sheet["B1"] = "Value"
        sheet["C1"] = "Outside"
        sheet["D4"] = "SecondArea"
        sheet["A2"] = ArrayFormula("A2:A2", "=SUM(B1:B1)")
        sheet["B2"] = DataTableFormula(ref="B2:B2", dt2D=False, r1="A1")
        sheet.print_area = ["A1:B2", "D4:E5"]
        book.save(source)
        book.close()

        first_inspection = inspect_source(source, ["Screen Layout"])
        second_inspection = inspect_source(source, ["Screen Layout"])
        self.assertEqual(first_inspection, second_inspection)
        by_ref = {record["source_ref"]: record for record in first_inspection["records"]}
        self.assertEqual(by_ref["Screen Layout!A2"]["kind"], "formula")
        self.assertEqual(by_ref["Screen Layout!B2"]["kind"], "formula")
        self.assertNotIn("0x", by_ref["Screen Layout!A2"]["value"])
        self.assertNotIn("0x", by_ref["Screen Layout!B2"]["value"])
        self.assertTrue(by_ref["Screen Layout!A1"]["inside_print_area"])
        self.assertFalse(by_ref["Screen Layout!C1"]["inside_print_area"])
        self.assertTrue(by_ref["Screen Layout!D4"]["inside_print_area"])

        mapping = {
            "schema_version": 2,
            "source_sha256": first_inspection["source_sha256"],
            "sheets": ["Screen Layout"],
            "tables": [
                {
                    "name": "Extracted",
                    "columns": ["label", "value"],
                    "rows": [{"label": ["Screen Layout!A1"], "value": ["Screen Layout!B1"]}],
                }
            ],
            "excluded": {"Screen Layout!C1": "Outside print area", "Screen Layout!D4": "Second area"},
        }
        destination = self.root / "formula-print-packet"
        report = build_packet(source, mapping, destination)
        self.assertEqual(report["status"], "draft_with_unresolved")
        verified = verify_packet(source, destination)
        self.assertEqual(verified["status"], "draft_with_unresolved")


def create_example(destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=False)
    source, mapping = create_generic_fixture(destination)
    (destination / "mapping.json").write_text(json.dumps(mapping, indent=2) + "\n", encoding="utf-8")
    build_packet(source, mapping, destination / "packet")
    report = verify_packet(source, destination / "packet")
    print("EXAMPLE_VERIFIED", json.dumps(report))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--example-dir", type=Path)
    parser.add_argument("--verify-example-dir", type=Path)
    arguments = parser.parse_args()
    if arguments.example_dir is not None:
        create_example(arguments.example_dir)
    elif arguments.verify_example_dir is not None:
        directory = arguments.verify_example_dir
        report = verify_packet(directory / "synthetic-report.xlsx", directory / "packet")
        print("EXAMPLE_VERIFIED", json.dumps(report))
    else:
        result = unittest.TextTestRunner(verbosity=1).run(
            unittest.defaultTestLoader.loadTestsFromTestCase(ReconstructionTests)
        )
        if result.wasSuccessful():
            print("RECONSTRUCTION_TESTS_OK")
        raise SystemExit(0 if result.wasSuccessful() else 1)
