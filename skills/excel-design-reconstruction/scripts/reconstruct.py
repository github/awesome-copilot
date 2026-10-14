from __future__ import annotations

import argparse
import hashlib
import io
import json
import posixpath
import re
import warnings
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree
from xml.sax import SAXException, handler, make_parser
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.table import Table, TableStyleInfo


FIELD_COLUMNS = ["field_id", "name", "data_type", "length", "required", "default", "condition", "notes"]
RULE_COLUMNS = ["rule_id", "condition", "action", "message", "notes"]
INVENTORY_COLUMNS = ["source_ref", "kind", "value", "source_type", "number_format", "merged_range", "hidden_row", "hidden_column", "inside_print_area", "state", "status", "reason"]
MAP_COLUMNS = ["target_table", "row_key", "column", "source_ref", "source_value"]
UNRESOLVED_COLUMNS = ["source_ref", "kind", "value", "reason"]
DEFAULT_FONT_NAME = "Noto Sans JP"
MAX_STORED_CELLS_PER_SHEET = 1_000_000
RELATIONSHIP_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
WORKSHEET_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
CHARTSHEET_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartsheet"
PREFLIGHT_RELATIONSHIP_TYPES = {
    f"{RELATIONSHIP_BASE}{name}"
    for name in (
        "chart", "chartsheet", "comments", "core-properties", "custom-properties",
        "drawing", "extended-properties", "officeDocument", "pivotCacheDefinition",
        "pivotCacheRecords", "pivotTable", "queryTable", "sharedStrings", "styles",
        "table", "theme", "vmlDrawing", "worksheet",
    )
}
CORE_PACKAGE_PARTS = {
    "[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml",
    "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml",
    "xl/sharedStrings.xml", "xl/theme/theme1.xml", "xl/calcChain.xml",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


class _SaxDoctypeGuard(handler.ContentHandler, handler.LexicalHandler):
    def startDTD(self, name, public_id, system_id) -> None:
        raise SAXException("DOCTYPE declarations are not supported")

    def endDTD(self) -> None:
        pass

    def startEntity(self, name) -> None:
        pass

    def endEntity(self, name) -> None:
        pass

    def startCDATA(self) -> None:
        pass

    def endCDATA(self) -> None:
        pass

    def comment(self, text) -> None:
        pass


def _new_sax_parser(content_handler: _SaxDoctypeGuard):
    parser = make_parser()
    parser.setFeature(handler.feature_namespaces, True)
    parser.setFeature(handler.feature_external_ges, False)
    parser.setFeature(handler.feature_external_pes, False)
    parser.setContentHandler(content_handler)
    parser.setProperty(handler.property_lexical_handler, content_handler)
    return parser


def _reject_doctype(xml_content: bytes) -> None:
    _new_sax_parser(_SaxDoctypeGuard()).parse(io.BytesIO(xml_content))


def _worksheet_paths(archive: ZipFile) -> dict[str, str]:
    workbook_xml = archive.read("xl/workbook.xml")
    relationships_xml = archive.read("xl/_rels/workbook.xml.rels")
    _reject_doctype(workbook_xml)
    _reject_doctype(relationships_xml)
    workbook = ElementTree.fromstring(workbook_xml)
    relationships = ElementTree.fromstring(relationships_xml)
    targets = {}
    for item in relationships:
        if _local_name(item.tag) != "Relationship":
            continue
        if not {"Id", "Type", "Target"} <= item.attrib.keys():
            raise ValueError("invalid workbook relationship")
        relationship_id = item.attrib["Id"]
        if relationship_id in targets:
            raise ValueError(f"duplicate workbook relationship id: {relationship_id}")
        targets[relationship_id] = item.attrib
    selected = {}
    for item in workbook.iter():
        if _local_name(item.tag) != "sheet":
            continue
        relation_id = next((value for key, value in item.attrib.items() if _local_name(key) == "id"), None)
        relationship = targets.get(relation_id)
        if relationship is None:
            raise ValueError(f"worksheet relationship not found: {item.attrib.get('name', '')}")
        relationship_type = relationship["Type"]
        if relationship_type == CHARTSHEET_RELATIONSHIP_TYPE:
            continue
        if relationship_type != WORKSHEET_RELATIONSHIP_TYPE:
            raise ValueError(f"unsupported sheet relationship: {item.attrib.get('name', '')}")
        target = relationship["Target"]
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        if not path.startswith("xl/worksheets/"):
            raise ValueError(f"unsupported worksheet package path: {path}")
        if path not in archive.namelist():
            raise ValueError(f"invalid worksheet package path: {path}")
        selected[item.attrib["name"]] = path
    return selected


def _relationship_source_path(relationships_path: str) -> str | None:
    if relationships_path == "_rels/.rels":
        return ""
    if "/_rels/" not in relationships_path or not relationships_path.endswith(".rels"):
        return None
    directory, filename = relationships_path.split("/_rels/", 1)
    return posixpath.join(directory, filename[:-5])


def _related_xml_parts(archive: ZipFile) -> set[str]:
    members = set(archive.namelist())
    related_parts = set()
    for relationships_path in members:
        if not relationships_path.endswith(".rels"):
            continue
        source_path = _relationship_source_path(relationships_path)
        if source_path is None:
            continue
        content = archive.read(relationships_path)
        _reject_doctype(content)
        relationships = ElementTree.fromstring(content)
        for relationship in relationships:
            if _local_name(relationship.tag) != "Relationship":
                continue
            if relationship.attrib.get("Type") not in PREFLIGHT_RELATIONSHIP_TYPES:
                continue
            if relationship.attrib.get("TargetMode") == "External":
                continue
            target = relationship.attrib.get("Target")
            if not target:
                continue
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
                posixpath.join(posixpath.dirname(source_path), target)
            )
            if path == ".." or path.startswith("../"):
                raise ValueError("relationship target escapes package")
            if path in members and path.lower().endswith((".xml", ".rels", ".vml")):
                related_parts.add(path)
    return related_parts


class _WorksheetCellCounter(_SaxDoctypeGuard):
    def __init__(self, sheet_name: str) -> None:
        super().__init__()
        self.sheet_name = sheet_name
        self.cell_count = 0

    def startElementNS(self, name, qname, attributes) -> None:
        if name[1] == "c":
            self.cell_count += 1
            if self.cell_count > MAX_STORED_CELLS_PER_SHEET:
                raise ValueError(f"sheet exceeds one million stored cells: {self.sheet_name}")


def _preflight_xml_parts(archive: ZipFile, worksheet_paths: dict[str, str]) -> None:
    worksheet_names = {path: name for name, path in worksheet_paths.items()}
    related_xml_parts = _related_xml_parts(archive)
    for member in archive.infolist():
        if not member.filename.lower().endswith((".xml", ".rels", ".vml")):
            continue
        if not (
            member.filename in CORE_PACKAGE_PARTS
            or member.filename.endswith((".rels", ".vml"))
            or member.filename.startswith((
                "xl/charts/", "xl/comments", "xl/drawings/", "xl/pivotCache/",
                "xl/pivotTables/", "xl/tables/",
            ))
            or member.filename in worksheet_names
            or member.filename in related_xml_parts
        ):
            continue
        content_handler = (
            _WorksheetCellCounter(worksheet_names[member.filename])
            if member.filename in worksheet_names else _SaxDoctypeGuard()
        )
        with archive.open(member) as stream:
            _new_sax_parser(content_handler).parse(stream)


def mapping_sha256(mapping: dict) -> str:
    canonical = json.dumps(mapping, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def approved_font_name(mapping: dict, requested_font_name: str | None = None) -> str:
    font_name = mapping.get("font_name", DEFAULT_FONT_NAME)
    if not isinstance(font_name, str) or not font_name.strip() or font_name != font_name.strip():
        raise ValueError("mapping font name must be a nonempty trimmed string")
    validate_text(font_name)
    if requested_font_name is not None:
        if not isinstance(requested_font_name, str) or requested_font_name.strip() != font_name:
            raise ValueError("font name must match approved mapping")
    return font_name


def _require_approved_mapping(mapping: dict, approved_mapping_sha256: str) -> str:
    actual = mapping_sha256(mapping)
    if not re.fullmatch(r"[0-9a-f]{64}", approved_mapping_sha256 or "") or actual != approved_mapping_sha256:
        raise ValueError("approved mapping hash mismatch")
    return actual


def inspect_source(source: Path, sheets: list[str]) -> dict:
    source = Path(source)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("a trusted local .xlsx file is required")
    if not sheets or len(set(sheets)) != len(sheets):
        raise ValueError("select a nonempty, unique sheet sequence")
    if source.stat().st_size > 256 * 1024 * 1024:
        raise ValueError("input exceeds 256 MiB")
    source_hash = sha256(source)
    records = []
    with ZipFile(source) as archive:
        if sum(member.file_size for member in archive.infolist()) > 1024 * 1024 * 1024:
            raise ValueError("expanded package exceeds 1 GiB")
        worksheet_paths = _worksheet_paths(archive)
        missing_sheets = sorted(set(sheets) - set(worksheet_paths))
        if missing_sheets:
            raise ValueError(f"unknown sheet: {', '.join(missing_sheets)}")
        _preflight_xml_parts(archive, worksheet_paths)
        stored_values = {}
        for sheet_name in sheets:
            worksheet = ElementTree.fromstring(archive.read(worksheet_paths[sheet_name]))
            namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for cell in worksheet.iter(f"{namespace}c"):
                cell_type = cell.get("t", "n")
                value = cell.find(f"{namespace}v")
                if cell_type in {"n", "d"} and value is not None and cell.find(f"{namespace}f") is None:
                    stored_values[(sheet_name, cell.get("r"))] = (value.text or "", cell_type)
        known_worksheet_paths = set(worksheet_paths.values())
        for name in sorted(archive.namelist()):
            if name.endswith("/"):
                continue
            if name not in CORE_PACKAGE_PARTS and name not in known_worksheet_paths:
                records.append({"source_ref": f"PART:{name}", "kind": "package_object",
                                "value": name, "source_type": "package", "number_format": "",
                                "merged_range": "", "hidden_row": None, "hidden_column": None,
                                "inside_print_area": None, "state": "workbook_scope"})
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        book = load_workbook(source, data_only=False, keep_links=False)
        try:
            for name in sheets:
                if name not in book.sheetnames:
                    raise ValueError(f"unknown sheet: {name}")
                sheet = book[name]
                merges = list(sheet.merged_cells.ranges)
                print_area = str(sheet.print_area or "")
                print_ranges = list(MultiCellRange(print_area)) if print_area else []
                for cell in sorted(sheet._cells.values(), key=lambda item: (item.row, item.column)):
                    if cell.value is None:
                        continue
                    hidden_column = any(
                        dimension.hidden and dimension.min <= cell.column <= dimension.max
                        for dimension in sheet.column_dimensions.values()
                    )
                    inside_print_area = (
                        any(cell.coordinate in cell_range for cell_range in print_ranges)
                        if print_ranges else None
                    )
                    stored_value, stored_type = stored_values.get(
                        (name, cell.coordinate), (str(cell.value), cell.data_type)
                    )
                    records.append({"source_ref": f"{name}!{cell.coordinate}",
                                    "kind": "formula" if cell.data_type == "f" else "cell",
                                    "value": stored_value, "source_type": stored_type,
                                    "number_format": str(cell.number_format),
                                    "merged_range": next((str(region) for region in merges if cell.coordinate in region), ""),
                                    "hidden_row": bool(sheet.row_dimensions[cell.row].hidden),
                                    "hidden_column": hidden_column,
                                    "inside_print_area": inside_print_area,
                                    "state": sheet.sheet_state})
                layout = {
                    "merged_ranges": [str(region) for region in merges],
                    "print_area": print_area,
                    "hidden_rows": [index for index, dimension in sheet.row_dimensions.items() if dimension.hidden],
                    "hidden_columns": [index for index, dimension in sheet.column_dimensions.items() if dimension.hidden],
                    "custom_row_height_count": sum(1 for item in sheet.row_dimensions.values() if item.height is not None),
                    "custom_column_width_count": sum(1 for item in sheet.column_dimensions.values() if item.width is not None),
                    "styled_cell_count": sum(1 for cell in sheet._cells.values() if cell.has_style),
                    "distinct_style_count": len({cell.style_id for cell in sheet._cells.values() if cell.has_style}),
                    "conditional_formatting_count": len(sheet.conditional_formatting),
                    "validation_count": len(sheet.data_validations.dataValidation),
                    "hyperlinks": [cell.coordinate for cell in sheet._cells.values() if getattr(cell, "hyperlink", None)],
                }
                records.append({"source_ref": f"LAYOUT:{name}", "kind": "layout_review",
                                "value": json.dumps(layout, ensure_ascii=False, sort_keys=True),
                                "source_type": "layout", "number_format": "",
                                "merged_range": "", "hidden_row": None, "hidden_column": None,
                                "inside_print_area": None, "state": sheet.sheet_state})
        finally:
            book.close()
        for index, warning in enumerate(caught, start=1):
            records.append({"source_ref": f"WARNING:{index}", "kind": "extraction_warning",
                            "value": str(warning.message), "source_type": "warning",
                            "number_format": "", "merged_range": "", "hidden_row": None,
                            "hidden_column": None, "inside_print_area": None, "state": "workbook_scope"})
    if sha256(source) != source_hash:
        raise ValueError("source hash changed during inspection")
    return {"schema_version": 1, "source_file": source.name, "source_sha256": source_hash,
            "sheets": sheets, "records": records}


def prepare_tables(inventory: dict, mapping: dict) -> dict:
    schema_version = mapping.get("schema_version")
    if schema_version == 1:
        if set(mapping) - {"schema_version", "source_sha256", "sheets", "fields", "rules", "excluded", "font_name"}:
            raise ValueError("unknown mapping keys")
        specifications = [
            {"name": "Fields", "columns": FIELD_COLUMNS, "rows": mapping.get("fields", [])},
            {"name": "Rules", "columns": RULE_COLUMNS, "rows": mapping.get("rules", [])},
        ]
    elif schema_version == 2:
        if set(mapping) - {"schema_version", "source_sha256", "sheets", "tables", "excluded", "font_name"}:
            raise ValueError("unknown mapping keys")
        specifications = mapping.get("tables")
        if not isinstance(specifications, list) or not specifications:
            raise ValueError("tables must be a nonempty array")
    else:
        raise ValueError("unsupported mapping schema")
    if mapping.get("source_sha256") != inventory["source_sha256"]:
        raise ValueError("mapping source hash mismatch")
    by_reference = {record["source_ref"]: record for record in inventory["records"]}
    used = set()
    source_map = []
    tables = {}
    reserved_names = {"history", "sourcemap", "sourceinventory", "unresolved"}
    for specification in specifications:
        if schema_version == 1:
            name = specification["name"]
            columns = specification["columns"]
            entries = specification["rows"]
        else:
            if not isinstance(specification, dict) or set(specification) != {"name", "columns", "rows"}:
                raise ValueError("each table needs name, columns, and rows")
            name = specification["name"]
            columns = specification["columns"]
            entries = specification["rows"]
        if (not isinstance(name, str) or not name.strip() or len(name) > 31
            or name.startswith("'") or name.endswith("'")
            or re.search(r"[:\\/?*\[\]]", name) or name.casefold() in reserved_names):
            raise ValueError("invalid or reserved table name")
        normalized_name = name.casefold()
        if normalized_name in {table_name.casefold() for table_name in tables}:
            raise ValueError("duplicate table name")
        if (not isinstance(columns, list) or not columns
                or not all(isinstance(column, str) and column.strip() for column in columns)):
            raise ValueError("table columns must be nonempty names")
        if len(columns) + 1 > 16384:
            raise ValueError("table exceeds Excel column limit including row_key")
        normalized_columns = [column.casefold() for column in columns]
        if "row_key" in normalized_columns or len(set(normalized_columns)) != len(columns):
            raise ValueError("table columns must be unique and cannot use row_key")
        if not isinstance(entries, list):
            raise ValueError("table rows must be an array")
        rows = []
        for index, entry in enumerate(entries, start=1):
            if not isinstance(entry, dict) or not entry or set(entry) - set(columns):
                raise ValueError("unknown or empty table row column")
            row_key = f"{name}-{index:04d}"
            row = [row_key]
            for column in columns:
                references = entry.get(column, [])
                if not isinstance(references, list) or (column in entry and not references):
                    raise ValueError("each mapped column requires a nonempty reference array")
                if not all(isinstance(reference, str) for reference in references) or len(set(references)) != len(references):
                    raise ValueError("duplicate or invalid reference")
                values = []
                for reference in references:
                    record = by_reference.get(reference)
                    if record is None:
                        raise ValueError(f"unknown reference or non-anchor cell: {reference}")
                    if record["kind"] == "formula":
                        raise ValueError("formula results cannot be inferred from stored formulas")
                    if record["kind"] != "cell":
                        raise ValueError("only stored cell values can populate semantic columns")
                    values.append(record["value"])
                    used.add(reference)
                    source_map.append([name, row_key, column, reference, record["value"]])
                row.append("\n".join(values))
            rows.append(row)
        tables[name] = {"columns": ["row_key", *columns], "rows": rows}
    excluded = mapping.get("excluded", {})
    if not isinstance(excluded, dict):
        raise ValueError("excluded must be an object with reasons")
    for reference, reason in excluded.items():
        if reference not in by_reference or reference in used:
            raise ValueError("invalid or mapped exclusion reference")
        if not isinstance(reason, str) or not reason.strip():
            raise ValueError("each exclusion needs a nonempty reason")
    inventory_rows = []
    unresolved = []
    for record in inventory["records"]:
        reference = record["source_ref"]
        status = "mapped" if reference in used else "excluded" if reference in excluded else "unresolved"
        reason = excluded.get(reference, "")
        if status == "unresolved":
            reason = "Needs human mapping or review; no semantic inference performed."
            unresolved.append([reference, record["kind"], record["value"], reason])
        inventory_rows.append([
            str(record[column]) if record[column] is not None else ""
            for column in INVENTORY_COLUMNS[:-2]
        ] + [status, reason])
    tables["SourceMap"] = {"columns": MAP_COLUMNS, "rows": source_map}
    tables["SourceInventory"] = {"columns": INVENTORY_COLUMNS, "rows": inventory_rows}
    tables["Unresolved"] = {"columns": UNRESOLVED_COLUMNS, "rows": unresolved}
    return tables


def validate_text(value: str) -> None:
    if len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise ValueError("value cannot be represented losslessly in an Excel cell")


def _create_output_workbook(tables: dict, font_name: str) -> Workbook:
    book = Workbook()
    book.remove(book.active)
    for table_index, (name, specification) in enumerate(tables.items(), start=1):
        sheet = book.create_sheet(name)
        sheet.append(specification["columns"])
        rows = specification["rows"] or [[""] * len(specification["columns"])]
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.data_type = "s"
                cell.number_format = "@"
                cell.font = Font(name=font_name, size=11)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column_index in range(1, len(specification["columns"]) + 1):
            sheet.column_dimensions[get_column_letter(column_index)].width = 24
        sheet.freeze_panes = "A2"
        table = Table(displayName=f"DataTable{table_index:04d}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    return book


def _verify_output_package(output: Path, tables: dict, font_name: str) -> None:
    canonical = io.BytesIO()
    book = _create_output_workbook(tables, font_name)
    try:
        book.save(canonical)
    finally:
        book.close()
    with ZipFile(canonical) as expected, ZipFile(output) as actual:
        members = actual.namelist()
        if len(members) != len(set(members)) or set(members) != set(expected.namelist()):
            raise ValueError("output package membership mismatch")
        for name in sorted(expected.namelist(), key=lambda part: (part != "xl/styles.xml", part)):
            expected_content = expected.read(name)
            detail = ""
            if name.startswith("xl/worksheets/"):
                detail = "; output table values differ from source mapping"
            elif name == "xl/styles.xml":
                detail = "; output font mismatch"
            if actual.getinfo(name).file_size > len(expected_content) + 128:
                raise ValueError(f"output package content size mismatch: {name}{detail}")
            actual_content = actual.read(name)
            if name == "docProps/core.xml":
                normalized = []
                for content in (expected_content, actual_content):
                    _reject_doctype(content)
                    root = ElementTree.fromstring(content)
                    for child in list(root):
                        if child.tag in {"{http://purl.org/dc/terms/}created", "{http://purl.org/dc/terms/}modified"}:
                            child.text = ""
                    normalized.append(ElementTree.tostring(root))
                expected_content, actual_content = normalized
            if actual_content != expected_content:
                raise ValueError(f"output package content mismatch: {name}{detail}")


def build_packet(
    source: Path,
    mapping: dict,
    destination: Path,
    *,
    approved_mapping_sha256: str,
    approved_by: str,
    font_name: str | None = None,
) -> dict:
    source, destination = Path(source), Path(destination)
    if destination.exists():
        raise FileExistsError("output directory already exists")
    approved_hash = _require_approved_mapping(mapping, approved_mapping_sha256)
    font_name = approved_font_name(mapping, font_name)
    if not isinstance(approved_by, str) or not approved_by.strip():
        raise ValueError("approval record name is required")
    inventory = inspect_source(source, mapping.get("sheets", []))
    tables = prepare_tables(inventory, mapping)
    for table in tables.values():
        if len(table["rows"]) > 1_048_575:
            raise ValueError("table exceeds Excel row limit")
        for row in table["rows"]:
            for value in row:
                validate_text(value)
    book = _create_output_workbook(tables, font_name)
    try:
        destination.mkdir(parents=True, exist_ok=False)
        output = destination / "reconstructed.xlsx"
        book.save(output)
    finally:
        book.close()
    if sha256(source) != inventory["source_sha256"]:
        raise ValueError("source hash changed during reconstruction")
    write_json(destination / "mapping.json", mapping)
    write_json(destination / "inventory.json", inventory)
    unresolved_count = len(tables["Unresolved"]["rows"])
    manifest = {
        "schema_version": 1, "source_file": source.name,
        "source_sha256": inventory["source_sha256"],
        "approval": {
            "approved_by": approved_by.strip(),
            "approved_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "mapping_sha256": approved_hash,
        },
        "font_name": font_name,
        "status": "draft_with_unresolved" if unresolved_count else "mapped_pending_semantic_review",
        "unresolved_count": unresolved_count,
        "row_counts": {name: len(table["rows"]) for name, table in tables.items()},
        "artifact_hashes": {name: sha256(destination / name) for name in ["reconstructed.xlsx", "mapping.json", "inventory.json"]},
    }
    write_json(destination / "manifest.json", manifest)
    return verify_packet(source, destination, approved_mapping_sha256=approved_hash)


def verify_packet(
    source: Path,
    destination: Path,
    *,
    approved_mapping_sha256: str,
) -> dict:
    source, destination = Path(source), Path(destination)
    manifest = json.loads((destination / "manifest.json").read_text(encoding="utf-8"))
    expected_artifacts = {"reconstructed.xlsx", "mapping.json", "inventory.json"}
    if set(manifest["artifact_hashes"]) != expected_artifacts:
        raise ValueError("unexpected artifact hash membership")
    for name, expected_hash in manifest["artifact_hashes"].items():
        if sha256(destination / name) != expected_hash:
            raise ValueError(f"artifact hash mismatch: {name}")
    mapping = json.loads((destination / "mapping.json").read_text(encoding="utf-8"))
    approved_hash = _require_approved_mapping(mapping, approved_mapping_sha256)
    font_name = approved_font_name(mapping)
    if manifest.get("font_name", DEFAULT_FONT_NAME) != font_name:
        raise ValueError("manifest font name mismatch")
    approval = manifest.get("approval", {})
    if (approval.get("mapping_sha256") != approved_hash
            or not isinstance(approval.get("approved_by"), str)
            or not approval["approved_by"].strip()):
        raise ValueError("approved mapping hash mismatch")
    try:
        datetime.fromisoformat(approval["approved_at"])
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError("approval timestamp is invalid") from error
    inventory = inspect_source(source, mapping.get("sheets", []))
    if inventory != json.loads((destination / "inventory.json").read_text(encoding="utf-8")) or inventory["source_sha256"] != manifest["source_sha256"]:
        raise ValueError("inventory or source hash mismatch")
    tables = prepare_tables(inventory, mapping)
    _verify_output_package(destination / "reconstructed.xlsx", tables, font_name)
    book = load_workbook(destination / "reconstructed.xlsx", data_only=False)
    try:
        if book.sheetnames != list(tables):
            raise ValueError("worksheet membership mismatch")
        for name, table in tables.items():
            sheet = book[name]
            if sheet.merged_cells.ranges or len(sheet.tables) != 1:
                raise ValueError("output must contain unmerged Excel tables")
            if any(cell.font.name != font_name for row in sheet for cell in row):
                raise ValueError("output font mismatch")
            expected_rows = table["rows"] or [[""] * len(table["columns"])]
            expected = [table["columns"], *expected_rows]
            actual = [[str(cell.value) if cell.value is not None else "" for cell in row] for row in sheet]
            if actual != expected or any(cell.data_type == "f" for row in sheet for cell in row):
                raise ValueError("output table values differ from source mapping")
            only_table = next(iter(sheet.tables.values()))
            if only_table.ref != sheet.dimensions:
                raise ValueError("Excel table does not cover output rows")
    finally:
        book.close()
    count = len(tables["Unresolved"]["rows"])
    status = "draft_with_unresolved" if count else "mapped_pending_semantic_review"
    if manifest["unresolved_count"] != count or manifest["status"] != status or manifest["row_counts"] != {name: len(table["rows"]) for name, table in tables.items()}:
        raise ValueError("manifest summary mismatch")
    return {
        "status": status,
        "unresolved_count": count,
        "row_counts": manifest["row_counts"],
        "approval": approval,
        "font_name": font_name,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Reconstruct layout-heavy Excel documents using explicit source mappings.")
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect")
    inspect.add_argument("--workbook", type=Path, required=True)
    inspect.add_argument("--sheet", action="append", required=True)
    inspect.add_argument("--out-dir", type=Path, required=True)
    build = commands.add_parser("build")
    build.add_argument("--workbook", type=Path, required=True)
    build.add_argument("--mapping", type=Path, required=True)
    build.add_argument("--approved-mapping-sha256", required=True)
    build.add_argument("--approved-by", required=True)
    build.add_argument("--font-name")
    build.add_argument("--out-dir", type=Path, required=True)
    fingerprint = commands.add_parser("fingerprint")
    fingerprint.add_argument("--workbook", type=Path, required=True)
    fingerprint.add_argument("--mapping", type=Path, required=True)
    verify = commands.add_parser("verify")
    verify.add_argument("--workbook", type=Path, required=True)
    verify.add_argument("--packet", type=Path, required=True)
    verify.add_argument("--approved-mapping-sha256", required=True)
    arguments = parser.parse_args()
    try:
        if arguments.command == "inspect":
            if arguments.out_dir.exists():
                raise FileExistsError("inspection directory already exists")
            inventory = inspect_source(arguments.workbook, arguments.sheet)
            arguments.out_dir.mkdir(parents=True, exist_ok=False)
            write_json(arguments.out_dir / "inventory.json", inventory)
            print("SOURCE_INSPECTED", len(inventory["records"]))
        elif arguments.command == "build":
            mapping = json.loads(arguments.mapping.read_text(encoding="utf-8"))
            report = build_packet(
                arguments.workbook, mapping, arguments.out_dir,
                approved_mapping_sha256=arguments.approved_mapping_sha256,
                approved_by=arguments.approved_by,
                font_name=arguments.font_name,
            )
            print("RECONSTRUCTION_VERIFIED", json.dumps(report))
        elif arguments.command == "fingerprint":
            mapping = json.loads(arguments.mapping.read_text(encoding="utf-8"))
            inventory = inspect_source(arguments.workbook, mapping.get("sheets", []))
            tables = prepare_tables(inventory, mapping)
            data_tables = {
                name: len(specification["rows"])
                for name, specification in tables.items()
                if name not in {"SourceMap", "SourceInventory", "Unresolved"}
            }
            print("MAPPING_REVIEW_READY", mapping_sha256(mapping), json.dumps({
                "font_name": approved_font_name(mapping),
                "tables": data_tables,
                "unresolved": len(tables["Unresolved"]["rows"]),
            }))
        else:
            report = verify_packet(
                arguments.workbook, arguments.packet,
                approved_mapping_sha256=arguments.approved_mapping_sha256,
            )
            print("RECONSTRUCTION_VERIFIED", json.dumps(report))
        return 0
    except Exception as error:
        print(f"RECONSTRUCTION_ERROR {type(error).__name__}: {error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
