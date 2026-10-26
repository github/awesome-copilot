from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from lxml import etree
from markdown_it import MarkdownIt


SVG_NAMESPACE = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXTENSION_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
IMAGE_RELATIONSHIP = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
RELATIONSHIPS_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/content-types"
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DRAWING_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
WORD_DRAWING_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
RELATIONSHIP_ATTRIBUTE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
SVG_CONFIG = {
    "theme": "default",
    "htmlLabels": False,
    "themeVariables": {"fontFamily": "Noto Sans JP, sans-serif"},
    "flowchart": {"htmlLabels": False},
}
PNG_FALLBACK = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\x0bIDAT"
    b"x\x9cc\x00\x02\x00\x00\x05\x00\x01\xa5\xf6E\x8c\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _resolve_cli(cli_name: str) -> str:
    resolved = shutil.which(cli_name)
    if resolved:
        return resolved
    candidate = Path(cli_name)
    if candidate.is_file():
        return str(candidate.resolve())
    raise FileNotFoundError(
        f"Mermaid CLI '{cli_name}' was not found. Install @mermaid-js/mermaid-cli or pass --mermaid-cli."
    )


def _command_prefix(cli: str) -> list[str]:
    cli_path = Path(cli)
    if os.name != "nt" or cli_path.suffix.lower() not in {".cmd", ".bat"}:
        return [cli]
    node = shutil.which("node")
    if node is None:
        raise FileNotFoundError("Node.js is required to invoke Mermaid CLI without a shell on Windows.")
    for directory in cli_path.resolve().parents:
        candidates = (
            directory / "@mermaid-js/mermaid-cli/src/cli.js",
            directory / "node_modules/@mermaid-js/mermaid-cli/src/cli.js",
        )
        entry_point = next((candidate for candidate in candidates if candidate.is_file()), None)
        if entry_point is not None:
            return [node, str(entry_point)]
    raise FileNotFoundError(
        "Could not locate the Mermaid CLI JavaScript entry point beside the Windows command shim."
    )


def _puppeteer_config(directory: Path) -> list[str]:
    browser = os.environ.get("MERMAID_BROWSER_PATH")
    candidates = [Path(browser)] if browser else []
    if os.name == "nt":
        candidates.extend(
            [
                Path(os.environ.get("PROGRAMFILES", "C:/Program Files"))
                / "Google/Chrome/Application/chrome.exe",
                Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)"))
                / "Microsoft/Edge/Application/msedge.exe",
            ]
        )
    browser_path = next((path.resolve() for path in candidates if path.is_file()), None)
    if browser_path is None:
        return []
    config_path = directory / "puppeteer-config.json"
    config_path.write_text(json.dumps({"executablePath": str(browser_path)}), encoding="utf-8")
    return ["-p", str(config_path)]


def _run_command(command: list[str], *, timeout: float = 120.0) -> None:
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
    except subprocess.TimeoutExpired as error:
        raise RuntimeError(f"Mermaid rendering timed out after {timeout:g}s.") from error
    if result.returncode:
        details = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError(f"Mermaid rendering failed (exit {result.returncode}): {details}")


def _render_mermaid(
    code: str, diagram_number: int, temp_root: Path, cli: str,
    font_family: str = "Noto Sans JP",
    *,
    timeout: float = 120.0,
) -> tuple[Path, Path, float, float]:
    diagram_dir = temp_root / f"diagram-{diagram_number}"
    diagram_dir.mkdir()
    source_path = diagram_dir / "diagram.mmd"
    config_path = diagram_dir / "mermaid-config.json"
    svg_path = diagram_dir / "diagram.svg"
    png_path = diagram_dir / "diagram.png"
    source_path.write_text(code.rstrip() + "\n", encoding="utf-8")
    config = {**SVG_CONFIG, "themeVariables": {"fontFamily": f"{font_family}, sans-serif"}}
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")

    command_prefix = _command_prefix(cli)
    common = ["-i", str(source_path), "-c", str(config_path), *_puppeteer_config(diagram_dir)]
    _run_command([*command_prefix, *common, "-o", str(svg_path), "-e", "svg"], timeout=timeout)
    _run_command([*command_prefix, *common, "-o", str(png_path), "-e", "png", "-s", "3"], timeout=timeout)

    if not svg_path.is_file() or not png_path.is_file():
        raise RuntimeError("Mermaid CLI completed without producing both SVG and PNG outputs.")
    width, height = _validate_svg(svg_path)
    return svg_path, png_path, width, height


def _svg_length(value: str | None) -> float | None:
    if not value:
        return None
    match = re.fullmatch(r"\s*([0-9]*\.?[0-9]+)(?:px)?\s*", value)
    return float(match.group(1)) if match else None


def _validate_svg(svg_path: Path) -> tuple[float, float]:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, recover=False)
    try:
        root = etree.parse(str(svg_path), parser).getroot()
    except etree.XMLSyntaxError as error:
        raise RuntimeError(f"Mermaid CLI produced invalid SVG: {error}") from error

    if etree.QName(root).localname != "svg":
        raise RuntimeError("Mermaid CLI output does not have an SVG root element.")

    view_box = root.get("viewBox")
    if view_box:
        values = [float(value) for value in re.split(r"[ ,]+", view_box.strip())]
        if len(values) == 4 and values[2] > 0 and values[3] > 0:
            return values[2], values[3]

    width = _svg_length(root.get("width"))
    height = _svg_length(root.get("height"))
    if width and height:
        return width, height
    raise RuntimeError("Mermaid SVG has no usable viewBox or width/height dimensions.")


def _add_inline_content(paragraph: Any, tokens: list[Any], *, header: bool = False) -> None:
    emphasis: list[str] = []
    link_stack: list[str] = []
    for token in tokens:
        if token.type in {"strong_open", "em_open", "s_open"}:
            emphasis.append(token.type.removesuffix("_open"))
        elif token.type in {"strong_close", "em_close", "s_close"}:
            if emphasis:
                emphasis.pop()
        elif token.type == "link_open":
            link_stack.append(token.attrGet("href") or "")
        elif token.type == "link_close":
            href = link_stack.pop() if link_stack else ""
            if href:
                paragraph.add_run(f" ({href})")
        elif token.type in {"text", "code_inline", "html_inline"}:
            text = token.content
            if token.type == "html_inline":
                text = re.sub(r"<[^>]*>", "", text)
            run = paragraph.add_run(text)
            run.bold = header or "strong" in emphasis
            run.italic = "em" in emphasis
            run.font.strike = "s" in emphasis
            if token.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "image":
            paragraph.add_run(token.content or "[image]")
            source = token.attrGet("src")
            if source:
                paragraph.add_run(f" ({source})")


def _new_paragraph(document: Any, list_stack: list[dict[str, Any]], quote_depth: int) -> Any:
    if quote_depth:
        return document.add_paragraph(style="Quote")
    if list_stack:
        depth = min(len(list_stack), 3)
        suffix = f" {depth}" if depth > 1 else ""
        current_list = list_stack[-1]
        if current_list.get("item_continuation", False):
            return document.add_paragraph(style=f"List Continue{suffix}")
        current_list["item_continuation"] = True
        if current_list["ordered"]:
            paragraph = document.add_paragraph(style=f"List Continue{suffix}")
            run = paragraph.add_run(f"{current_list['next_number']}. ")
            run.bold = False
            current_list["next_number"] += 1
            return paragraph
        return document.add_paragraph(style=f"List Bullet{suffix}")
    return document.add_paragraph()


def _diagram_description(code: str, diagram_number: int, current_heading: str | None) -> tuple[str, str]:
    acc_title = re.search(r"(?m)^\s*accTitle\s*:\s*(.+?)\s*$", code)
    acc_descr = re.search(r"(?m)^\s*accDescr\s*:\s*(.+?)\s*$", code)
    first_statement = next(
        (
            line.strip()
            for line in code.splitlines()
            if line.strip() and not line.strip().startswith("%%")
        ),
        "diagram",
    )
    diagram_kind = first_statement.split()[0]
    title = (
        acc_title.group(1)
        if acc_title
        else f"{current_heading} — Mermaid diagram {diagram_number}"
        if current_heading
        else f"Mermaid diagram {diagram_number}"
    )
    if acc_descr:
        description = f"{title}: {acc_descr.group(1)} ({diagram_kind}; vector SVG)."
    else:
        description = f"{title} ({diagram_kind}; vector SVG)."
    return title, description


def _add_code_block(document: Any, content: str) -> None:
    paragraph = document.add_paragraph(style="No Spacing")
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(content.rstrip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _add_diagram(
    document: Any,
    svg_path: Path,
    png_path: Path,
    diagram_number: int,
    svg_width: float,
    svg_height: float,
    *,
    code: str = "",
    current_heading: str | None = None,
) -> None:
    section = document.sections[-1]
    available_width = section.page_width - section.left_margin - section.right_margin
    available_height = section.page_height - section.top_margin - section.bottom_margin
    width = float(available_width)
    height = width * svg_height / svg_width
    if height >= available_height * 0.70:
        scale = min(1.0, float(available_height) / height)
        width *= scale
        height *= scale
    if height >= available_height * 0.70:
        print(
            f"Warning: Mermaid diagram {diagram_number} uses most of a page; verify label readability in Word.",
            file=sys.stderr,
        )

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    if height >= available_height * 0.70:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run()
    picture = run.add_picture(str(png_path), width=Emu(round(width)), height=Emu(round(height)))
    title, description = _diagram_description(code, diagram_number, current_heading)
    picture._inline.docPr.set("name", f"mermaid-svg-{diagram_number}")
    picture._inline.docPr.set("title", title)
    picture._inline.docPr.set("descr", description)


def _render_table(document: Any, tokens: list[Any], start: int) -> int:
    rows: list[tuple[list[Any], bool]] = []
    current_cells: list[Any] | None = None
    in_header = False
    index = start + 1
    while index < len(tokens) and tokens[index].type != "table_close":
        token = tokens[index]
        if token.type == "thead_open":
            in_header = True
        elif token.type == "tbody_open":
            in_header = False
        elif token.type == "tr_open":
            current_cells = []
        elif token.type == "inline" and current_cells is not None:
            current_cells.append(token)
        elif token.type == "tr_close" and current_cells is not None:
            rows.append((current_cells, in_header))
            current_cells = None
        index += 1

    column_count = max((len(cells) for cells, _ in rows), default=0)
    if column_count:
        table = document.add_table(rows=len(rows), cols=column_count)
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        for row_index, (cells, is_header) in enumerate(rows):
            for column_index, inline_token in enumerate(cells):
                cell = table.cell(row_index, column_index)
                _add_inline_content(cell.paragraphs[0], inline_token.children or [], header=is_header)
    return index + 1


def _convert_blocks(
    document: Any, markdown_text: str, temp_root: Path, cli_name: str,
    font_family: str = "Noto Sans JP",
    *,
    timeout: float = 120.0,
) -> int:
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(markdown_text)
    list_stack: list[dict[str, Any]] = []
    quote_depth = 0
    current_paragraph = None
    in_heading = False
    current_heading: str | None = None
    diagram_number = 0
    cli: str | None = None
    index = 0

    while index < len(tokens):
        token = tokens[index]
        token_type = token.type

        if token_type == "table_open":
            index = _render_table(document, tokens, index)
            current_paragraph = None
            continue
        if token_type in {"bullet_list_open", "ordered_list_open"}:
            start_attr = token.attrGet("start") if token_type == "ordered_list_open" else None
            start_number = int(start_attr) if start_attr and str(start_attr).isdigit() else 1
            list_stack.append(
                {
                    "ordered": token_type == "ordered_list_open",
                    "next_number": start_number,
                    "item_continuation": False,
                }
            )
        elif token_type == "list_item_open":
            if list_stack:
                list_stack[-1]["item_continuation"] = False
        elif token_type == "list_item_close":
            if list_stack:
                list_stack[-1]["item_continuation"] = False
        elif token_type in {"bullet_list_close", "ordered_list_close"}:
            if list_stack:
                list_stack.pop()
        elif token_type == "blockquote_open":
            quote_depth += 1
        elif token_type == "blockquote_close":
            quote_depth = max(0, quote_depth - 1)
        elif token_type == "paragraph_open":
            current_paragraph = _new_paragraph(document, list_stack, quote_depth)
        elif token_type == "heading_open":
            level = int(token.tag[1])
            current_paragraph = document.add_paragraph(style=f"Heading {level}")
            current_paragraph.paragraph_format.keep_with_next = True
            in_heading = True
        elif token_type == "inline":
            if current_paragraph is None:
                current_paragraph = _new_paragraph(document, list_stack, quote_depth)
            _add_inline_content(current_paragraph, token.children or [])
            if in_heading and token.content.strip():
                current_heading = token.content.strip()
        elif token_type == "paragraph_close":
            current_paragraph = None
        elif token_type == "heading_close":
            current_paragraph = None
            in_heading = False
        elif token_type == "fence":
            if token.info.strip().split(maxsplit=1)[0:1] == ["mermaid"]:
                diagram_number += 1
                if cli is None:
                    cli = _resolve_cli(cli_name)
                svg_path, png_path, svg_width, svg_height = _render_mermaid(
                    token.content, diagram_number, temp_root, cli, font_family, timeout=timeout
                )
                _add_diagram(
                    document,
                    svg_path,
                    png_path,
                    diagram_number,
                    svg_width,
                    svg_height,
                    code=token.content,
                    current_heading=current_heading,
                )
            else:
                _add_code_block(document, token.content)
            current_paragraph = None
        elif token_type == "code_block":
            _add_code_block(document, token.content)
            current_paragraph = None
        elif token_type == "html_block":
            plain_text = re.sub(r"<[^>]*>", "", token.content).strip()
            if plain_text:
                document.add_paragraph(plain_text)
            current_paragraph = None
        elif token_type == "hr":
            document.add_paragraph()
            current_paragraph = None
        index += 1

    return diagram_number


def _svg_relationship_id(rels_root: Any) -> str:
    numbers = [
        int(match.group(1))
        for relationship in rels_root.findall(f"{{{RELATIONSHIPS_NAMESPACE}}}Relationship")
        if (match := re.fullmatch(r"rId([0-9]+)", relationship.get("Id", "")))
    ]
    return f"rId{max(numbers, default=0) + 1}"


def _inject_svg_parts(
    docx_path: Path, svg_paths: list[Path], *, temp_root: Path | None = None
) -> None:
    if not svg_paths:
        return

    ns = {
        "a": DRAWING_NAMESPACE,
        "asvg": SVG_NAMESPACE,
        "wp": WORD_DRAWING_NAMESPACE,
    }
    temporary_docx = (
        temp_root / "injected-output.docx"
        if temp_root is not None
        else docx_path.with_name(f"{docx_path.stem}.tmp{docx_path.suffix}")
    )
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
        temporary_docx, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        document_xml = etree.fromstring(source.read("word/document.xml"))
        relationships_xml = etree.fromstring(source.read("word/_rels/document.xml.rels"))
        content_types_xml = etree.fromstring(source.read("[Content_Types].xml"))
        svg_default = content_types_xml.find(f"{{{CONTENT_TYPES_NAMESPACE}}}Default[@Extension='svg']")
        if svg_default is None:
            svg_default = etree.SubElement(content_types_xml, f"{{{CONTENT_TYPES_NAMESPACE}}}Default")
            svg_default.set("Extension", "svg")
            svg_default.set("ContentType", "image/svg+xml")

        replacements = {
            "word/document.xml": None,
            "word/_rels/document.xml.rels": None,
            "[Content_Types].xml": None,
        }
        for number, svg_path in enumerate(svg_paths, start=1):
            marker = f"mermaid-svg-{number}"
            inline = document_xml.xpath(
                f".//wp:inline[wp:docPr[@name='{marker}']]", namespaces=ns
            )
            if len(inline) != 1:
                raise RuntimeError(f"Could not locate the Word drawing for Mermaid diagram {number}.")
            blip = inline[0].find(".//a:blip", namespaces=ns)
            if blip is None:
                raise RuntimeError(f"Word drawing for Mermaid diagram {number} has no image reference.")

            relationship_id = _svg_relationship_id(relationships_xml)
            relationship = etree.SubElement(
                relationships_xml, f"{{{RELATIONSHIPS_NAMESPACE}}}Relationship"
            )
            relationship.set("Id", relationship_id)
            relationship.set("Type", IMAGE_RELATIONSHIP)
            relationship.set("Target", f"media/mermaid-{number}.svg")

            extension_list = blip.find("a:extLst", namespaces=ns)
            if extension_list is None:
                extension_list = etree.SubElement(blip, f"{{{DRAWING_NAMESPACE}}}extLst")
            extension = etree.SubElement(extension_list, f"{{{DRAWING_NAMESPACE}}}ext")
            extension.set("uri", SVG_EXTENSION_URI)
            svg_blip = etree.SubElement(extension, f"{{{SVG_NAMESPACE}}}svgBlip")
            svg_blip.set(RELATIONSHIP_ATTRIBUTE, relationship_id)
            replacements[f"word/media/mermaid-{number}.svg"] = svg_path.read_bytes()

        replacements["word/document.xml"] = etree.tostring(
            document_xml, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        replacements["word/_rels/document.xml.rels"] = etree.tostring(
            relationships_xml, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        replacements["[Content_Types].xml"] = etree.tostring(
            content_types_xml, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        for item in source.infolist():
            if item.filename not in replacements:
                target.writestr(item, source.read(item.filename))
        for name, data in replacements.items():
            if data is not None:
                target.writestr(name, data)
    temporary_docx.replace(docx_path)


def _extract_frontmatter(markdown_text: str) -> tuple[str | None, str]:
    lines = markdown_text.splitlines()
    if not lines or lines[0].strip() != "---":
        return None, markdown_text
    closing_index = next(
        (index for index in range(1, len(lines)) if lines[index].strip() in {"---", "..."}),
        None,
    )
    if closing_index is None:
        return None, markdown_text

    metadata_lines = lines[1:closing_index]
    first_metadata_line = next((line.strip() for line in metadata_lines if line.strip()), "")
    if not re.match(r"^[A-Za-z_][A-Za-z0-9_-]*\s*:", first_metadata_line):
        return None, markdown_text

    frontmatter = "\n".join(metadata_lines)
    title_match = re.search(r"(?m)^title:\s*(.*?)\s*$", frontmatter)
    title = title_match.group(1).strip("\"'") if title_match else None
    body = "\n".join(lines[closing_index + 1 :])
    if body:
        body += "\n"
    return title or None, body


def convert_markdown(
    input_path: Path, output_path: Path, cli_name: str,
    *, font_family: str = "Noto Sans JP", timeout: float = 120.0,
) -> None:
    input_path = input_path.resolve()
    output_path = output_path.resolve()
    if input_path == output_path:
        raise ValueError("Input Markdown and output DOCX paths must be different.")
    if input_path.suffix.lower() not in {".md", ".markdown"}:
        raise ValueError("Input file must use the .md or .markdown extension.")
    if output_path.suffix.lower() != ".docx":
        raise ValueError("Output file must use the .docx extension.")
    if output_path.exists():
        raise FileExistsError(f"Output file already exists; refusing to overwrite it: {output_path}")
    font_family = font_family.strip()
    if not font_family:
        raise ValueError("Font family must not be empty.")
    if timeout <= 0:
        raise ValueError("Timeout must be greater than zero.")

    title, markdown_text = _extract_frontmatter(input_path.read_text(encoding="utf-8-sig"))
    document = Document()
    for style in document.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.font.name = font_family
            style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_family)
    document.core_properties.title = title or input_path.stem
    first_heading = re.search(r"(?m)^#\s+(.+?)\s*$", markdown_text)
    if title and (first_heading is None or first_heading.group(1) != title):
        document.add_heading(title, level=0)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(
        prefix=".markdown-to-word-mermaid-", dir=output_path.parent
    ) as temporary_directory:
        temp_root = Path(temporary_directory)
        staged_docx = temp_root / "staged-output.docx"
        diagram_count = _convert_blocks(
            document, markdown_text, temp_root, cli_name, font_family, timeout=timeout
        )
        document.save(staged_docx)
        svg_paths = [temp_root / f"diagram-{number}" / "diagram.svg" for number in range(1, diagram_count + 1)]
        _inject_svg_parts(staged_docx, svg_paths, temp_root=temp_root)
        _publish_docx_exclusively(staged_docx, output_path)


def _publish_docx_exclusively(staged_docx: Path, output_path: Path) -> None:
    try:
        os.link(staged_docx, output_path)
    except FileExistsError as error:
        raise FileExistsError(
            f"Output file already exists; refusing to overwrite it: {output_path}"
        ) from error
    except OSError as error:
        if os.name != "nt":
            raise OSError(
                f"Cannot publish DOCX atomically because hard links are unavailable: {output_path}"
            ) from error
        try:
            os.rename(staged_docx, output_path)
        except FileExistsError as conflict:
            raise FileExistsError(
                f"Output file already exists; refusing to overwrite it: {output_path}"
            ) from conflict
    else:
        staged_docx.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert Markdown to DOCX and embed Mermaid diagrams as SVG vectors."
    )
    parser.add_argument("input", type=Path, help="Input Markdown file")
    parser.add_argument("--output", type=Path, help="Output DOCX path")
    parser.add_argument("--mermaid-cli", default="mmdc", help="Mermaid CLI executable (default: mmdc)")
    parser.add_argument("--font-family", default="Noto Sans JP", help="Installed font for text and diagrams (default: Noto Sans JP)")
    parser.add_argument("--timeout", type=float, default=120.0, help="Per-command Mermaid CLI timeout in seconds (default: 120)")
    args = parser.parse_args()
    output_path = args.output or args.input.with_suffix(".docx")

    try:
        convert_markdown(
            args.input,
            output_path,
            args.mermaid_cli,
            font_family=args.font_family,
            timeout=args.timeout,
        )
    except (FileNotFoundError, OSError, RuntimeError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(output_path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
