from __future__ import annotations

import argparse
import importlib
import json
import sys
import tempfile
import unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from capture_visual_evidence import capture_selected_sheets, verify_visual_evidence
import capture_visual_evidence
from excel_review_evidence import inspect_workbook, write_inspection


class VisualCaptureSmokeTests(unittest.TestCase):
    evidence_dir: Path | None = None

    @staticmethod
    def create_failure_fixture(root: Path, multiple_pages: bool = False) -> Path:
        source = root / "failure-fixture.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Layout"
        for row_index in range(1, 121 if multiple_pages else 5):
            sheet.cell(row_index, 1, f"Row {row_index}")
            sheet.row_dimensions[row_index].height = 14
        sheet.print_area = "A1:D120" if multiple_pages else "A1:D4"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0 if multiple_pages else 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        workbook.save(source)
        workbook.close()
        return source

    def test_live_hidden_sheet_requires_explicit_permission(self) -> None:
        for state in ["hidden", "veryHidden"]:
            with self.subTest(state=state), tempfile.TemporaryDirectory() as temp_dir:
                root = Path(temp_dir)
                source = root / "hidden-fixture.xlsx"
                workbook = Workbook()
                workbook.active.title = "Fields"
                layout = workbook.create_sheet("Layout")
                layout["A1"] = "Hidden fixture"
                layout.print_area = "A1:D4"
                layout.sheet_state = state
                workbook.save(source)
                workbook.close()
                structure = inspect_workbook(source)
                next(sheet for sheet in structure["sheets"] if sheet["name"] == "Layout")["state"] = "visible"
                denied = root / "denied"
                with self.assertRaisesRegex(RuntimeError, "did not complete"):
                    capture_selected_sheets(source, structure, ["Layout"], denied)
                manifest = json.loads((denied / "visual-evidence-manifest.json").read_text(encoding="utf-8"))
                self.assertEqual(manifest["status"], "partial_failure")
                self.assertIn("include_hidden=True", manifest["selected_sheets"][0]["error"])
                self.assertEqual(list(denied.rglob("*.pdf")), [])
                self.assertEqual(list(denied.rglob("*.png")), [])
                self.assertEqual(manifest["source_sha256"], manifest["source_sha256_after"])
                self.assertEqual(inspect_workbook(source)["sheets"][1]["state"], state)
                approved = root / "approved"
                result = capture_selected_sheets(source, structure, ["Layout"], approved, include_hidden=True)
                self.assertEqual(result["status"], "complete")
                verify_visual_evidence(source, approved / "visual-evidence-manifest.json", ["Layout"])
                self.assertEqual(inspect_workbook(source)["sheets"][1]["state"], state)

    def test_unlisted_artifacts_outside_sheets_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = self.create_failure_fixture(root)
            destination = root / "evidence"
            manifest = capture_selected_sheets(source, inspect_workbook(source), ["Layout"], destination)
            manifest_path = destination / "visual-evidence-manifest.json"
            pdf = destination / manifest["selected_sheets"][0]["pdf_file"]
            image = destination / manifest["selected_sheets"][0]["pages"][0]["image_file"]
            for relative_path in ["extra.pdf", "extra.png", "nested/extra.pdf", "nested/extra.png"]:
                with self.subTest(path=relative_path):
                    extra = destination / relative_path
                    extra.parent.mkdir(parents=True, exist_ok=True)
                    extra.write_bytes((pdf if extra.suffix == ".pdf" else image).read_bytes())
                    try:
                        with self.assertRaisesRegex(ValueError, "exactly match"):
                            verify_visual_evidence(source, manifest_path, ["Layout"])
                    finally:
                        extra.unlink()
                    verify_visual_evidence(source, manifest_path, ["Layout"])

    def test_verification_failure_is_persisted_as_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = self.create_failure_fixture(root)
            destination = root / "evidence"
            with patch.object(capture_visual_evidence, "verify_visual_evidence", side_effect=ValueError("injected mismatch")):
                with self.assertRaisesRegex(ValueError, "injected mismatch"):
                    capture_selected_sheets(source, inspect_workbook(source), ["Layout"], destination)
            manifest = json.loads((destination / "visual-evidence-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "error")
            self.assertIn("verification failed", manifest["error"])
            self.assertEqual(len(list(destination.rglob("*.pdf"))), 1)
            self.assertEqual(len(list(destination.rglob("*.png"))), 1)
            with self.assertRaisesRegex(ValueError, "not complete"):
                verify_visual_evidence(source, destination / "visual-evidence-manifest.json", ["Layout"])

    def test_second_page_failure_retains_partial_artifact_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = self.create_failure_fixture(root, multiple_pages=True)
            destination = root / "evidence"
            pdfium = importlib.import_module("pypdfium2")
            original_render = pdfium.PdfPage.render
            render_calls = 0

            def render_or_fail(page, *args, **kwargs):
                nonlocal render_calls
                render_calls += 1
                if render_calls == 2:
                    raise RuntimeError("injected second-page failure")
                return original_render(page, *args, **kwargs)

            with patch.object(pdfium.PdfPage, "render", render_or_fail):
                with self.assertRaisesRegex(RuntimeError, "did not complete"):
                    capture_selected_sheets(source, inspect_workbook(source), ["Layout"], destination)
            manifest = json.loads((destination / "visual-evidence-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "partial_failure")
            record = manifest["selected_sheets"][0]
            self.assertEqual(record["status"], "error")
            self.assertIn("injected second-page failure", record["error"])
            self.assertGreaterEqual(record["page_count"], 2)
            self.assertTrue(record["capture_basis"])
            self.assertEqual(len(record["pdf_sha256"]), 64)
            self.assertTrue((destination / record["pdf_file"]).is_file())
            self.assertEqual(len(record["pages"]), 1)
            self.assertTrue((destination / record["pages"][0]["image_file"]).is_file())
            self.assertEqual(len(record["pages"][0]["image_sha256"]), 64)
            self.assertEqual(len(list(destination.rglob("*.png"))), 1)
            self.assertEqual(manifest["source_sha256"], manifest["source_sha256_after"])

    def test_save_failure_retains_written_image_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = self.create_failure_fixture(root)
            destination = root / "evidence"
            pillow_image = importlib.import_module("PIL.Image")
            original_save = pillow_image.Image.save

            def save_then_fail(image, *args, **kwargs):
                original_save(image, *args, **kwargs)
                raise RuntimeError("injected failure after image save")

            with patch.object(pillow_image.Image, "save", save_then_fail):
                with self.assertRaisesRegex(RuntimeError, "did not complete"):
                    capture_selected_sheets(source, inspect_workbook(source), ["Layout"], destination)
            manifest = json.loads((destination / "visual-evidence-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "partial_failure")
            record = manifest["selected_sheets"][0]
            self.assertIn("injected failure after image save", record["error"])
            self.assertEqual(len(record["pages"]), 1)
            page = record["pages"][0]
            self.assertEqual(page["page"], 1)
            self.assertTrue((destination / page["image_file"]).is_file())
            self.assertEqual(page["image_sha256"], "")
            self.assertEqual(len(list(destination.rglob("*.png"))), 1)

    def test_japanese_graph_paper_preserves_shapes_and_labels(self) -> None:
        directory = (
            tempfile.TemporaryDirectory() if self.evidence_dir is None
            else nullcontext(str(self.evidence_dir))
        )
        with directory as temp_dir:
            root = Path(temp_dir) / "graph-paper"
            root.mkdir(parents=True, exist_ok=False)
            source_path = root / "japanese-graph-paper.xlsx"
            sheet_name = "\u753b\u9762\u8a2d\u8a08"
            title = "\u7533\u8acb\u5165\u529b\u753b\u9762"
            applicant = "\u7533\u8acb\u8005\u540d"
            amount = "\u7533\u8acb\u91d1\u984d"
            shape_labels = ["\u5165\u529b", "\u78ba\u8a8d", "\u767b\u9332", "\u4fee\u6b63"]
            workbook = Workbook()
            workbook.active.title = "Fields"
            workbook.active.append(["ID", "Label"])
            workbook.active.append(["FLD-001", applicant])
            layout = workbook.create_sheet(sheet_name)
            for column_index in range(1, 31):
                layout.column_dimensions[get_column_letter(column_index)].width = 2.4
            for row_index in range(1, 41):
                layout.row_dimensions[row_index].height = 12
            border_side = Side(style="thin", color="444444")
            regions = [
                ("B2:AD4", title, "DDEBF7"),
                ("B7:J9", applicant, "E2EFDA"),
                ("K7:AD9", "FLD-001", "FFFFFF"),
                ("B11:J13", amount, "E2EFDA"),
                ("K11:AD13", "100,000", "FFFFFF"),
                ("B16:AD17", "\u51e6\u7406\u30d5\u30ed\u30fc", "DDEBF7"),
                ("R20:T21", "\u53ef", "FFFFFF"),
                ("Q27:S28", "\u5426", "FFFFFF"),
            ]
            for cell_range, label, fill_color in regions:
                for row in layout[cell_range]:
                    for cell in row:
                        cell.border = Border(
                            left=border_side, right=border_side,
                            top=border_side, bottom=border_side,
                        )
                        cell.fill = PatternFill("solid", fgColor=fill_color)
                layout.merge_cells(cell_range)
                cell = layout[cell_range.split(":")[0]]
                cell.value = label
                cell.font = Font(name="Noto Sans JP", size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            layout.print_area = "A1:AD40"
            layout.sheet_view.showGridLines = False
            layout.page_setup.paperSize = layout.PAPERSIZE_A4
            layout.page_setup.fitToWidth = 1
            layout.page_setup.fitToHeight = 1
            layout.sheet_properties.pageSetUpPr.fitToPage = True
            workbook.save(source_path)
            workbook.close()

            pythoncom = importlib.import_module("pythoncom")
            win32_client = importlib.import_module("win32com.client")
            pythoncom.CoInitialize()
            application = None
            drawing_workbook = None
            try:
                application = win32_client.DispatchEx("Excel.Application")
                application.Visible = False
                application.DisplayAlerts = False
                application.EnableEvents = False
                application.AutomationSecurity = 3
                drawing_workbook = application.Workbooks.Open(
                    Filename=str(source_path.resolve()), UpdateLinks=0, AddToMru=False,
                )
                sheet = drawing_workbook.Worksheets(sheet_name)
                self.assertLess(sheet.Columns(1).ColumnWidth, 3)
                self.assertAlmostEqual(sheet.Columns(1).ColumnWidth, sheet.Columns(30).ColumnWidth)
                self.assertEqual(sheet.Rows(1).RowHeight, 12)
                self.assertEqual(sheet.Rows(40).RowHeight, 12)
                shape_regions = ["B20:G23", "K19:Q25", "U20:AD23", "K31:Q34"]
                shapes = []
                for shape_index, (cell_range, label) in enumerate(zip(shape_regions, shape_labels)):
                    region = sheet.Range(cell_range)
                    shape = sheet.Shapes.AddShape(
                        4 if shape_index == 1 else 1,
                        region.Left, region.Top, region.Width, region.Height,
                    )
                    shape.TextFrame2.TextRange.Text = label
                    shape.TextFrame2.TextRange.Font.Name = "Noto Sans JP"
                    shape.TextFrame2.TextRange.Font.Size = 11
                    shape.TextFrame2.TextRange.Font.Fill.ForeColor.RGB = 0
                    shape.TextFrame2.TextRange.ParagraphFormat.Alignment = 2
                    shape.TextFrame2.VerticalAnchor = 3
                    shape.Fill.ForeColor.RGB = 0xFFFFFF
                    shape.Line.ForeColor.RGB = 0
                    shapes.append(shape)
                start, decision, finish, correction = shapes
                arrows = [
                    (start.Left + start.Width, start.Top + start.Height / 2,
                     decision.Left, decision.Top + decision.Height / 2),
                    (decision.Left + decision.Width, decision.Top + decision.Height / 2,
                     finish.Left, finish.Top + finish.Height / 2),
                    (decision.Left + decision.Width / 2, decision.Top + decision.Height,
                     correction.Left + correction.Width / 2, correction.Top),
                ]
                for start_x, start_y, end_x, end_y in arrows:
                    connector = sheet.Shapes.AddConnector(1, start_x, start_y, end_x, end_y)
                    connector.Line.ForeColor.RGB = 0
                    connector.Line.Weight = 1.5
                    connector.Line.EndArrowheadStyle = 3
                self.assertEqual(sheet.Shapes.Count, 7)
                drawing_workbook.Save()
            finally:
                try:
                    if drawing_workbook is not None:
                        drawing_workbook.Close(SaveChanges=False)
                finally:
                    try:
                        if application is not None:
                            application.Quit()
                    finally:
                        pythoncom.CoUninitialize()

            packet = write_inspection(source_path, root / "inspection")
            structure = packet["structure"]
            metadata = next(sheet for sheet in structure["sheets"] if sheet["name"] == sheet_name)
            self.assertEqual(metadata["drawing_anchor_count"], 7)
            self.assertGreater(metadata["custom_column_width_count"], 0)
            self.assertEqual(metadata["custom_row_height_count"], 40)
            self.assertIn("B2:AD4", metadata["merged_ranges"])
            candidate = next(
                item for item in packet["visual_candidates"]["targets"]
                if item["sheet_name"] == sheet_name
            )
            self.assertIn("drawing_objects", candidate["reasons"])
            cell_values = [cell["value"] for cell in metadata["cells"]]
            self.assertIn(title, cell_values)
            for label in shape_labels:
                self.assertNotIn(label, cell_values)
            output_dir = root / "visual-evidence"
            manifest = capture_selected_sheets(
                source_path, structure, [sheet_name], output_dir,
                selection_note="Synthetic Japanese graph-paper layout and flow arrows",
            )
            self.assertEqual(manifest["source_sha256"], manifest["source_sha256_after"])
            self.assertEqual([record["sheet_name"] for record in manifest["selected_sheets"]], [sheet_name])
            self.assertEqual(len(list(output_dir.rglob("*.pdf"))), 1)
            self.assertEqual(len(list(output_dir.rglob("*.png"))), 1)
            record = manifest["selected_sheets"][0]
            self.assertEqual(record["page_count"], 1)
            pdfium = importlib.import_module("pypdfium2")
            with pdfium.PdfDocument(output_dir / record["pdf_file"]) as document:
                page = document[0]
                try:
                    text_page = page.get_textpage()
                    try:
                        rendered_text = "".join(text_page.get_text_bounded().split())
                    finally:
                        text_page.close()
                finally:
                    page.close()
                for label in [title, applicant, amount, *shape_labels, "\u53ef", "\u5426", "100,000"]:
                    self.assertIn(label, rendered_text)
            verify_visual_evidence(source_path, output_dir / "visual-evidence-manifest.json", [sheet_name])
            print("JAPANESE_GRAPH_PAPER_CAPTURE_OK")

    def test_captures_only_explicitly_selected_sheet_without_overwriting(self) -> None:
        directory = (
            tempfile.TemporaryDirectory() if self.evidence_dir is None
            else nullcontext(str(self.evidence_dir))
        )
        with directory as temp_dir:
            root = Path(temp_dir)
            root.mkdir(parents=True, exist_ok=True)
            source_path = root / "synthetic-review.xlsx"
            workbook = Workbook()
            workbook.active.title = "Fields"
            workbook.active.append(["ID", "Label"])
            workbook.active.append(["FLD-001", "Applicant"])
            layout = workbook.create_sheet("Layout")
            layout.merge_cells("A1:C1")
            layout["A1"] = "Form title"
            layout["A3"] = "Amount"
            layout["B3"] = 100
            layout.print_area = "A1:C6"
            workbook.save(source_path)
            workbook.close()

            structure = inspect_workbook(source_path)
            output_dir = root / "visual-evidence"
            manifest = capture_selected_sheets(
                source_path,
                structure,
                ["Layout"],
                output_dir,
            )

            self.assertEqual(manifest["status"], "complete")
            self.assertEqual(
                [record["sheet_name"] for record in manifest["selected_sheets"]],
                ["Layout"],
            )
            self.assertEqual(len(list(output_dir.rglob("*.pdf"))), 1)
            self.assertGreaterEqual(len(list(output_dir.rglob("*.png"))), 1)
            self.assertTrue((output_dir / "visual-evidence-manifest.json").is_file())
            saved_manifest = json.loads(
                (output_dir / "visual-evidence-manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(saved_manifest["source_sha256"], structure["source_sha256"])
            verify_visual_evidence(source_path, output_dir / "visual-evidence-manifest.json", ["Layout"])
            with self.assertRaises(ValueError):
                verify_visual_evidence(source_path, output_dir / "visual-evidence-manifest.json", ["Fields"])
            image_path = next(output_dir.rglob("*.png"))
            original_image = image_path.read_bytes()
            with image_path.open("ab") as stream:
                stream.write(b"modified")
            with self.assertRaises(ValueError):
                verify_visual_evidence(source_path, output_dir / "visual-evidence-manifest.json", ["Layout"])
            image_path.write_bytes(original_image)
            verify_visual_evidence(source_path, output_dir / "visual-evidence-manifest.json", ["Layout"])

            with self.assertRaises(FileExistsError):
                capture_selected_sheets(source_path, structure, ["Layout"], output_dir)

            failed_dir = root / "failed-capture"
            with patch("win32com.client.DispatchEx", side_effect=RuntimeError("Excel unavailable")):
                with self.assertRaises(RuntimeError):
                    capture_selected_sheets(source_path, structure, ["Layout"], failed_dir)
            failed_manifest = json.loads(
                (failed_dir / "visual-evidence-manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(failed_manifest["status"], "error")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    if args.evidence_dir is not None and args.evidence_dir.exists():
        parser.error("evidence directory already exists")
    VisualCaptureSmokeTests.evidence_dir = args.evidence_dir
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(VisualCaptureSmokeTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if result.wasSuccessful():
        print("VISUAL_CAPTURE_SMOKE_OK")
    raise SystemExit(0 if result.wasSuccessful() else 1)