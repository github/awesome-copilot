from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from excel_review_evidence import inspect_workbook, suggest_visual_targets, write_inspection
import excel_review_evidence


class WorkbookPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "synthetic-design.xlsx"
        self._create_workbook(self.workbook_path)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    @staticmethod
    def _create_workbook(path: Path) -> None:
        workbook = Workbook()
        table = workbook.active
        table.title = "Fields"
        table.append(["ID", "Label", "Required"])
        table.append(["FLD-001", "Applicant", "Yes"])
        table["D2"] = "=COUNTA(A2:A2)"

        layout = workbook.create_sheet("Layout")
        layout.merge_cells("A1:C1")
        layout["A1"] = "Screen title"
        layout["A3"] = "Amount"
        layout["B3"] = 100
        layout.print_area = "A1:C8"
        layout.row_dimensions[9].hidden = True
        layout.append(["Category", "Value"])
        layout.append(["A", 1])
        layout.append(["B", 2])
        chart = BarChart()
        chart.add_data(Reference(layout, min_col=2, min_row=10, max_row=11), titles_from_data=False)
        chart.set_categories(Reference(layout, min_col=1, min_row=10, max_row=11))
        layout.add_chart(chart, "E2")

        hidden = workbook.create_sheet("Internal notes")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "not visible by default"

        workbook.save(path)
        workbook.close()

    def test_extracts_content_and_marks_visual_review_candidates(self) -> None:
        structure = inspect_workbook(self.workbook_path)

        self.assertEqual(structure["sheet_count"], 3)
        self.assertEqual(len(structure["source_sha256"]), 64)
        self.assertEqual(structure["sheets"][0]["name"], "Fields")
        self.assertEqual(structure["sheets"][0]["formula_count"], 1)
        self.assertEqual(structure["sheets"][1]["merged_ranges"], ["A1:C1"])
        self.assertEqual(structure["sheets"][1]["chart_count"], 1)
        self.assertEqual(structure["sheets"][1]["print_area"], "'Layout'!$A$1:$C$8")
        self.assertEqual(structure["sheets"][2]["state"], "hidden")

        targets = suggest_visual_targets(structure)
        by_name = {item["sheet_name"]: item["reasons"] for item in targets}
        self.assertNotIn("Fields", by_name)
        self.assertIn("Layout", by_name)
        self.assertIn("chart", by_name["Layout"])
        self.assertIn("merged_cells", by_name["Layout"])
        self.assertNotIn("Internal notes", by_name)

    def test_counts_physical_columns_in_hidden_dimension_spans(self) -> None:
        cases = [
            ([("A", "C", True)], 3),
            ([("A", "C", True), ("E", "E", True), ("G", "I", False)], 4),
            ([("A", "C", True), ("B", "D", True)], 4),
            ([("B", "B", True)], 1),
            ([("A", "C", False)], 0),
        ]
        for case_index, (spans, expected_count) in enumerate(cases):
            with self.subTest(spans=spans):
                path = Path(self.temp_dir.name) / f"hidden-columns-{case_index}.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Layout"
                sheet["A1"] = "Synthetic layout"
                for start, end, hidden in spans:
                    sheet.column_dimensions.group(start, end, hidden=hidden)
                workbook.save(path)
                workbook.close()
                original_hash = excel_review_evidence._sha256(path)
                structure = inspect_workbook(path)
                self.assertEqual(structure["sheets"][0]["hidden_column_count"], expected_count)
                self.assertEqual(excel_review_evidence._sha256(path), original_hash)
                targets = suggest_visual_targets(structure)
                reasons = {reason for target in targets for reason in target["reasons"]}
                self.assertEqual("hidden_rows_or_columns" in reasons, expected_count > 0)

    def test_rejects_non_workbook_input(self) -> None:
        invalid_path = Path(self.temp_dir.name) / "not-a-workbook.xlsx"
        invalid_path.write_text("not an xlsx archive", encoding="utf-8")

        with self.assertRaises(ValueError):
            inspect_workbook(invalid_path)

    def test_rejects_source_changes_during_inspection(self) -> None:
        original_inspector = excel_review_evidence._inspect_sheet

        def inspect_and_change(worksheet):
            result = original_inspector(worksheet)
            with self.workbook_path.open("ab") as source:
                source.write(b"\n")
            return result

        output = Path(self.temp_dir.name) / "changed-inspection"
        with patch.object(excel_review_evidence, "_inspect_sheet", side_effect=inspect_and_change):
            with self.assertRaisesRegex(ValueError, "changed during inspection"):
                write_inspection(self.workbook_path, output)
        self.assertFalse(output.exists())

    def test_writes_packet_and_rejects_overwrite_and_csv(self) -> None:
        output = Path(self.temp_dir.name) / "inspection"
        result = write_inspection(self.workbook_path, output)
        self.assertTrue((output / "workbook-structure.json").is_file())
        self.assertTrue((output / "visual-candidates.json").is_file())
        self.assertTrue(result["visual_candidates"]["human_selection_required"])
        with self.assertRaises(FileExistsError):
            write_inspection(self.workbook_path, output)
        csv_path = Path(self.temp_dir.name) / "fields.csv"
        csv_path.write_text("id,label\n1,Applicant", encoding="utf-8")
        with self.assertRaises(ValueError):
            inspect_workbook(csv_path)


if __name__ == "__main__":
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(WorkbookPipelineTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if result.wasSuccessful():
        print("WORKBOOK_PIPELINE_TESTS_OK")
    raise SystemExit(0 if result.wasSuccessful() else 1)