from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import warnings
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
MAX_ARCHIVE_BYTES = 256 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
MAX_NONEMPTY_CELLS = 1_000_000
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_workbook(path: Path) -> None:
    if not path.is_file():
        raise ValueError(f"workbook does not exist: {path.name}")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"unsupported workbook type: {path.suffix or '(no extension)'}")
    if path.stat().st_size > MAX_ARCHIVE_BYTES:
        raise ValueError(f"workbook exceeds the {MAX_ARCHIVE_BYTES}-byte input limit")
    if not is_zipfile(path):
        raise ValueError("input is not a valid XLSX/XLSM package")
    try:
        with ZipFile(path) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("workbook package is missing xl/workbook.xml")
            expanded_size = sum(member.file_size for member in archive.infolist())
            if expanded_size > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("workbook package exceeds the expanded-size limit")
            corrupt_member = archive.testzip()
            if corrupt_member:
                raise ValueError(f"workbook package contains a corrupt member: {corrupt_member}")
    except BadZipFile as error:
        raise ValueError("input is not a valid XLSX/XLSM package") from error


def inspect_workbook(source_path: str | Path) -> dict:
    """Extract cell and worksheet metadata without modifying the source workbook."""
    path = Path(source_path)
    _validate_workbook(path)
    source_sha256 = _sha256(path)

    extraction_warnings: list[str] = []
    drawing_metadata, drawing_warnings = _inspect_drawing_parts(path)
    try:
        with warnings.catch_warnings(record=True) as caught_warnings:
            warnings.simplefilter("always")
            workbook = load_workbook(
                path,
                read_only=False,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
            try:
                sheets = []
                for worksheet in workbook.worksheets:
                    sheet = _inspect_sheet(worksheet)
                    sheet.update(drawing_metadata.get(worksheet.title, {}))
                    sheets.append(sheet)
            finally:
                workbook.close()
        extraction_warnings = [str(item.message) for item in caught_warnings]
    except Exception as error:
        raise ValueError(f"could not inspect workbook: {type(error).__name__}: {error}") from error

    if _sha256(path) != source_sha256:
        raise ValueError("source workbook changed during inspection")

    return {
        "schema_version": 1,
        "source_file": path.name,
        "source_sha256": source_sha256,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sheet_count": len(sheets),
        "extraction_warnings": extraction_warnings + drawing_warnings,
        "sheets": sheets,
    }


def _inspect_sheet(worksheet) -> dict:
    cells = []
    formula_count = 0
    if len(worksheet._cells) > MAX_NONEMPTY_CELLS:
        raise ValueError(
            f"sheet {worksheet.title!r} exceeds the {MAX_NONEMPTY_CELLS}-cell inspection limit"
        )
    for cell in sorted(worksheet._cells.values(), key=lambda item: (item.row, item.column)):
        if cell.value is None:
            continue
        value = str(cell.value)
        if cell.data_type == "f":
            formula_count += 1
        cells.append({
            "coordinate": cell.coordinate,
            "value": value,
            "data_type": cell.data_type,
        })

    row_dimensions = list(worksheet.row_dimensions.values())
    column_dimensions = list(worksheet.column_dimensions.values())
    chart_count = len(worksheet._charts)
    image_count = len(worksheet._images)

    return {
        "name": worksheet.title,
        "state": worksheet.sheet_state,
        "max_row": worksheet.max_row or 0,
        "max_column": worksheet.max_column or 0,
        "nonempty_cell_count": len(cells),
        "formula_count": formula_count,
        "cells": cells,
        "merged_ranges": sorted(str(item) for item in worksheet.merged_cells.ranges),
        "freeze_panes": str(worksheet.freeze_panes or ""),
        "print_area": str(worksheet.print_area or ""),
        "page_orientation": worksheet.page_setup.orientation,
        "fit_to_width": worksheet.page_setup.fitToWidth,
        "fit_to_height": worksheet.page_setup.fitToHeight,
        "hidden_row_count": sum(1 for item in row_dimensions if item.hidden),
        "hidden_column_count": sum(1 for item in column_dimensions if item.hidden),
        "custom_row_height_count": sum(1 for item in row_dimensions if item.height is not None),
        "custom_column_width_count": sum(1 for item in column_dimensions if item.width is not None),
        "chart_count": chart_count,
        "image_count": image_count,
        "visual_object_count": chart_count + image_count,
        "drawing_anchor_count": 0,
        "legacy_drawing_count": 0,
    }


def _relationship_part(source_part: str) -> str:
    return posixpath.join(
        posixpath.dirname(source_part),
        "_rels",
        f"{posixpath.basename(source_part)}.rels",
    )


def _read_relationships(archive: ZipFile, source_part: str) -> dict[str, ElementTree.Element]:
    try:
        root = ElementTree.fromstring(archive.read(_relationship_part(source_part)))
    except KeyError:
        return {}
    return {
        item.attrib["Id"]: item
        for item in root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }


def _resolve_part(source_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _inspect_drawing_parts(path: Path) -> tuple[dict[str, dict], list[str]]:
    metadata: dict[str, dict] = {}
    extraction_warnings: list[str] = []
    workbook_part = "xl/workbook.xml"

    with ZipFile(path) as archive:
        workbook_root = ElementTree.fromstring(archive.read(workbook_part))
        workbook_relationships = _read_relationships(archive, workbook_part)
        sheets_element = workbook_root.find(f"{{{MAIN_NS}}}sheets")
        if sheets_element is None:
            raise ValueError("workbook has no sheets element")

        for sheet_element in sheets_element:
            sheet_name = sheet_element.get("name", "")
            relationship_id = sheet_element.get(f"{{{DOCUMENT_REL_NS}}}id", "")
            workbook_relationship = workbook_relationships.get(relationship_id)
            if workbook_relationship is None:
                extraction_warnings.append(f"worksheet relationship missing: {sheet_name}")
                continue

            worksheet_part = _resolve_part(workbook_part, workbook_relationship.get("Target", ""))
            worksheet_root = ElementTree.fromstring(archive.read(worksheet_part))
            worksheet_relationships = _read_relationships(archive, worksheet_part)
            drawing_anchor_count = 0
            legacy_drawing_count = 0

            for tag in ("drawing", "legacyDrawing", "legacyDrawingHF", "picture"):
                for reference in worksheet_root.findall(f".//{{{MAIN_NS}}}{tag}"):
                    reference_id = reference.get(f"{{{DOCUMENT_REL_NS}}}id", "")
                    relation = worksheet_relationships.get(reference_id)
                    if relation is None:
                        extraction_warnings.append(f"{tag} relationship missing: {sheet_name}")
                        continue
                    target_part = _resolve_part(worksheet_part, relation.get("Target", ""))
                    if tag != "drawing":
                        legacy_drawing_count += 1
                        continue
                    drawing_root = ElementTree.fromstring(archive.read(target_part))
                    drawing_anchor_count += sum(
                        1
                        for element in drawing_root.iter()
                        if element.tag in {
                            f"{{{DRAWING_NS}}}oneCellAnchor",
                            f"{{{DRAWING_NS}}}twoCellAnchor",
                            f"{{{DRAWING_NS}}}absoluteAnchor",
                        }
                    )

            metadata[sheet_name] = {
                "drawing_anchor_count": drawing_anchor_count,
                "legacy_drawing_count": legacy_drawing_count,
            }

    return metadata, extraction_warnings


def suggest_visual_targets(structure: dict) -> list[dict]:
    """Suggest visible sheets whose meaning may depend on layout or embedded visuals."""
    suggestions = []
    for sheet in structure.get("sheets", []):
        if sheet.get("state") != "visible":
            continue

        reasons = []
        if sheet.get("chart_count", 0):
            reasons.append("chart")
        if sheet.get("image_count", 0):
            reasons.append("embedded_images")
        if sheet.get("drawing_anchor_count", 0) > sheet.get("chart_count", 0) + sheet.get("image_count", 0):
            reasons.append("drawing_objects")
        if sheet.get("legacy_drawing_count", 0):
            reasons.append("legacy_drawing_objects")
        if sheet.get("merged_ranges"):
            reasons.append("merged_cells")
        if sheet.get("print_area"):
            reasons.append("print_area")
        if sheet.get("hidden_row_count", 0) or sheet.get("hidden_column_count", 0):
            reasons.append("hidden_rows_or_columns")
        if sheet.get("custom_row_height_count", 0) or sheet.get("custom_column_width_count", 0):
            reasons.append("custom_dimensions")

        if reasons:
            suggestions.append({"sheet_name": sheet["name"], "reasons": reasons})
    return suggestions


def _write_json(path: Path, value: dict) -> None:
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary_path.replace(path)


def write_inspection(source_path: str | Path, output_dir: str | Path) -> dict:
    source = Path(source_path)
    destination = Path(output_dir)
    if destination.exists():
        raise FileExistsError(f"output directory already exists: {destination}")

    structure = inspect_workbook(source)
    suggestions = {
        "schema_version": 1,
        "source_file": structure["source_file"],
        "source_sha256": structure["source_sha256"],
        "human_selection_required": True,
        "targets": suggest_visual_targets(structure),
    }

    destination.mkdir(parents=True, exist_ok=False)
    _write_json(destination / "workbook-structure.json", structure)
    _write_json(destination / "visual-candidates.json", suggestions)
    print(
        "WORKBOOK_STRUCTURE_OK "
        f"sheets={structure['sheet_count']} "
        f"visual_candidates={len(suggestions['targets'])} "
        f"output={destination}"
    )
    return {"structure": structure, "visual_candidates": suggestions}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract Excel workbook structure and suggest sheets for visual review."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)
    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--workbook", required=True, type=Path)
    inspect_parser.add_argument("--out-dir", required=True, type=Path)
    args = parser.parse_args()

    try:
        if args.command == "inspect":
            write_inspection(args.workbook, args.out_dir)
            return 0
    except Exception as error:
        print(f"WORKBOOK_STRUCTURE_ERROR {type(error).__name__}: {error}")
        return 1
    return 2


if __name__ == "__main__":
    raise SystemExit(main())